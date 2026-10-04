import streamlit as st
import json
import re
import os
import time
import tempfile
import urllib.request
import threading
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from datetime import datetime, date, timedelta
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from fpdf import FPDF
import streamlit.components.v1 as components

from db import (
    autenticar, crear_usuario, usuario_existe, contar_admins,
    obtener_vehiculos_de_usuario, crear_vehiculo, eliminar_vehiculo,
    placas_existen, obtener_espacios, obtener_registro_activo_de_usuario,
    generar_qr_imagen,
    obtener_vehiculo_por_placas, obtener_registro_activo_por_vehiculo,
    registrar_entrada, registrar_salida, obtener_vehiculos_dentro,
    decodificar_qr_de_imagen, imagen_a_base64, base64_a_bytes,
    obtener_todos_usuarios, eliminar_usuario, obtener_todos_los_registros,
    generar_siguiente_id, id_estudiante_existe,
    actualizar_usuario, id_estudiante_existe_otro,
    activar_usuario, desactivar_usuario,
    registrar_log, obtener_logs, obtener_acciones_unicas, contar_logs,
    obtener_historial_usuario, contar_visitas_usuario,
    desbloquear_usuario,
    cambiar_password_usuario,
    actualizar_telefono_usuario,
    limpiar_cache,
    obtener_registros_paginado, contar_registros_filtrados,
    contar_logs_filtrados,
    obtener_historial_vehiculo, buscar_vehiculos_admin,
    contar_usuarios_por_estado,
    enviar_mensaje, obtener_mensajes_para_usuario, contar_mensajes_no_leidos,
    marcar_mensaje_leido, obtener_mensajes_enviados, eliminar_mensaje,
    obtener_vehiculos_alerta,
    crear_tabla_mensajes, crear_indices,
    listar_todas_las_placas,
)

try:
    from streamlit_webrtc import webrtc_streamer, VideoProcessorBase, WebRtcMode
    import av
    import cv2
    WEBRTC_DISPONIBLE = True
except ImportError:
    WEBRTC_DISPONIBLE = False

LOGO_COMPLETO_URL = "https://raw.githubusercontent.com/cj3302718-netizen/Estacionamiento-Yahualica/main/logo_completo.png"
LOGO_ESCUDO_URL = "https://raw.githubusercontent.com/cj3302718-netizen/Estacionamiento-Yahualica/main/logo_escudo.png"

st.set_page_config(page_title="CUYPARK", page_icon=LOGO_ESCUDO_URL, layout="centered")

APP_VERSION = "v5.0-anti-none"
st.session_state["_app_version"] = APP_VERSION

COLOR_VINO = "#7B1B2E"
COLOR_VINO_CLARO = "#D7192D"
COLOR_DORADO = "#C9A961"
COLOR_CREMA = "#F5F0E8"
COLOR_AZUL = "#0066B3"

SEGUNDOS_AUTOCAPTURA = 5

st.markdown("""
<style>
    @keyframes fadeInUp { from { opacity: 0; transform: translateY(20px); } to { opacity: 1; transform: translateY(0); } }
    @keyframes fadeInScale { from { opacity: 0; transform: scale(0.95); } to { opacity: 1; transform: scale(1); } }
    @keyframes slideInRight { from { opacity: 0; transform: translateX(100px); } to { opacity: 1; transform: translateX(0); } }
    @keyframes pulse { 0%,100% { transform: scale(1); box-shadow: 0 0 30px rgba(201,169,97,.6),0 0 60px rgba(123,27,46,.4);} 50% { transform: scale(1.03); box-shadow: 0 0 40px rgba(201,169,97,.8),0 0 80px rgba(123,27,46,.6);} }
    @keyframes shimmer { 0% { background-position: -200% center; } 100% { background-position: 200% center; } }
    @keyframes pulseAlert { 0%,100% { opacity: 1; } 50% { opacity: .65; } }

    .stApp { background: radial-gradient(ellipse at top, #1a0a0f 0%, #0d0407 40%, #050203 100%) !important; }
    header[data-testid="stHeader"] { background: rgba(0,0,0,0) !important; }
    section[data-testid="stSidebar"] { background: #0f0508 !important; }

    .titulo-principal { text-align:center; color:#C9A961; font-size:2rem; font-weight:800; margin-bottom:0; letter-spacing:2px; animation: fadeInUp .8s ease-out; }
    .subtitulo { text-align:center; color:#A89968; font-size:.9rem; margin-top:-8px; margin-bottom:20px; animation: fadeInUp 1s ease-out .2s both; }
    .panel-header { background: linear-gradient(135deg, #7B1B2E 0%, #D7192D 100%); padding:14px 16px; border-radius:12px; color:#FFF; font-weight:700; margin-bottom:16px; font-size:1.05rem; line-height:1.3; box-shadow:0 0 20px rgba(123,27,46,.4); animation: fadeInUp .6s ease-out; }
    .contador-card { background:#1a0a0f; border:2px solid #C9A961; border-radius:16px; padding:14px; text-align:center; box-shadow:0 0 15px rgba(201,169,97,.35); margin-bottom:10px; animation: fadeInScale .6s ease-out; transition: transform .3s, box-shadow .3s; }
    .contador-card:hover { transform: translateY(-4px); box-shadow:0 0 30px rgba(201,169,97,.6); }
    .contador-card h4 { color:#A89968; font-size:.75rem; text-transform:uppercase; letter-spacing:1px; margin:0; }
    .contador-card .numero { font-size:1.6rem; font-weight:800; color:#C9A961; margin-top:4px; }
    .edit-form { background:#0f0508; border:2px dashed #C9A961; border-radius:12px; padding:16px; margin-top:10px; animation: fadeInUp .4s ease-out; }
    .hist-item { background:#1a0a0f; border-left:3px solid #C9A961; border-radius:8px; padding:12px 14px; margin-bottom:8px; animation: fadeInUp .4s ease-out; transition: transform .2s, background .2s; }
    .hist-item:hover { transform: translateX(4px); background:#240e15; }
    .hist-duracion { color:#A89968; font-size:.8rem; }
    .log-row { background:#1a0a0f; border-left:3px solid #C9A961; border-radius:8px; padding:10px 14px; margin-bottom:8px; animation: fadeInUp .4s ease-out; transition: transform .2s, background .2s; }
    .log-row:hover { transform: translateX(4px); background:#240e15; }
    .log-accion { font-weight:700; font-size:.85rem; }
    .log-fecha { color:#A89968; font-size:.75rem; }
    .log-detalle { color:#E8DFD0; font-size:.85rem; margin-top:4px; }
    .ocupacion-label { display:flex; justify-content:space-between; font-size:.88rem; margin-bottom:4px; margin-top:8px; }
    .ocupacion-label b { color:#C9A961; }
    .kpi-card { background: linear-gradient(135deg, #1a0a0f 0%, #2a1015 100%); border:1px solid rgba(201,169,97,.35); border-radius:14px; padding:14px 16px; text-align:center; animation: fadeInUp .6s ease-out; transition: transform .3s, border-color .3s, box-shadow .3s; }
    .kpi-card:hover { transform: translateY(-3px); border-color:#C9A961; box-shadow:0 0 20px rgba(201,169,97,.4); }
    .kpi-card .kpi-label { color:#A89968; font-size:.75rem; text-transform:uppercase; letter-spacing:.8px; }
    .kpi-card .kpi-value { font-size:1.7rem; font-weight:800; color:#F5F0E8; margin-top:2px; }
    .kpi-card .kpi-delta { font-size:.8rem; margin-top:2px; }
    .delta-up { color:#00ff88; } .delta-down { color:#ff6666; } .delta-neutral { color:#A89968; }
    .bienvenida-card { background: linear-gradient(135deg, #1a3a1a 0%, #2a5028 100%); border:2px solid #00ff88; border-radius:16px; padding:18px; text-align:center; margin-bottom:16px; box-shadow:0 0 20px rgba(0,255,136,.3); animation: fadeInScale .6s ease-out; }
    .bienvenida-card h3 { color:#00ff88; margin:0; font-size:1.1rem; }
    .bienvenida-card p { color:#fff; margin:6px 0 0 0; font-size:.95rem; }
    .bienvenida-card .placas { color:#00ff88; font-weight:800; font-size:1.3rem; }
    .val-ok { color:#00ff88; font-size:.8rem; margin-top:-8px; margin-bottom:8px; animation: fadeInUp .3s ease-out; }
    .val-error { color:#ff5555; font-size:.8rem; margin-top:-8px; margin-bottom:8px; animation: fadeInUp .3s ease-out; }
    .logo-medallon { display:inline-block; background: radial-gradient(circle at 30% 30%, #FFF 0%, #F5F0E8 60%, #E8DFD0 100%); border-radius:50%; padding:6px; border:3px solid #C9A961; box-shadow:0 0 30px rgba(201,169,97,.6), 0 0 60px rgba(123,27,46,.4); margin-bottom:16px; animation: pulse 3s ease-in-out infinite; }
    .logo-medallon img { width:130px; height:130px; border-radius:50%; display:block; object-fit:cover; }
    .logo-barra { display:inline-block; background: radial-gradient(circle, #FFF 0%, #F5F0E8 100%); border-radius:50%; padding:2px; border:2px solid #C9A961; box-shadow:0 0 8px rgba(201,169,97,.5); vertical-align:middle; margin-right:8px; transition: transform .3s, box-shadow .3s; }
    .logo-barra:hover { transform: scale(1.1) rotate(5deg); box-shadow:0 0 15px rgba(201,169,97,.8); }
    .logo-barra img { width:40px; height:40px; border-radius:50%; display:block; object-fit:cover; }
    .brand-cudy { display:flex; align-items:center; justify-content:flex-end; gap:10px; padding:6px 0; }
    .brand-cudy .texto { text-align:right; line-height:1.1; }
    .brand-cudy .texto .linea1 { color:#A89968; font-size:.65rem; letter-spacing:2px; text-transform:uppercase; font-weight:600; }
    .brand-cudy .texto .linea2 { color:#C9A961; font-size:.95rem; font-weight:800; letter-spacing:1.5px; }
    .stProgress > div > div > div > div { background: linear-gradient(90deg, #7B1B2E 0%, #C9A961 100%) !important; transition: width .6s ease; }
    .stButton > button { transition: all .25s ease !important; position: relative; overflow: hidden; }
    .stButton > button:hover { transform: translateY(-2px); box-shadow:0 6px 20px rgba(201,169,97,.4) !important; }
    .stButton > button:active { transform: translateY(0); }
    .stButton > button[kind="primary"] { background: linear-gradient(135deg, #7B1B2E 0%, #D7192D 100%) !important; border:none !important; }
    .stButton > button[kind="primary"]:hover { background: linear-gradient(135deg, #D7192D 0%, #7B1B2E 100%) !important; }
    .stTextInput > div > div > input:focus, .stTextArea > div > div > textarea:focus { border-color:#C9A961 !important; box-shadow:0 0 0 3px rgba(201,169,97,.2), 0 0 15px rgba(201,169,97,.4) !important; }
    div[data-testid="stToast"] { animation: slideInRight .4s ease-out !important; }
    div[data-testid="stDecoration"] { background: linear-gradient(90deg, #7B1B2E 0%, #C9A961 50%, #D7192D 100%) !important; height:3px !important; animation: shimmer 2s linear infinite; background-size:200% 100%; }
    .alert-badge { display:inline-block; padding:3px 10px; border-radius:12px; font-size:.72rem; font-weight:700; letter-spacing:.5px; animation: pulseAlert 2s ease-in-out infinite; }
    .alert-badge.critico { background: rgba(215,25,45,.2); color:#ff4444; border:1px solid #ff4444; box-shadow:0 0 10px rgba(215,25,45,.5); }
    .alert-badge.advertencia { background: rgba(255,170,0,.2); color:#ffaa00; border:1px solid #ffaa00; box-shadow:0 0 10px rgba(255,170,0,.4); }
    .alert-badge.ok { background: rgba(0,255,136,.15); color:#00ff88; border:1px solid #00ff88; }
    .alert-card { background: linear-gradient(135deg, #2a0a0f 0%, #1a0505 100%); border:2px solid #ff4444; border-radius:12px; padding:12px 14px; margin-bottom:10px; animation: fadeInScale .5s ease-out; box-shadow:0 0 15px rgba(215,25,45,.3); }
    .countdown-num { font-size:7rem; font-weight:900; text-align:center; color:#ff4444; line-height:1; animation: fadeInScale .3s ease-out; text-shadow: 0 0 40px rgba(255,68,68,.8); }
    .instruccion-grande { text-align:center; font-size:1.4rem; font-weight:800; color:#C9A961; padding:16px; background: linear-gradient(135deg, rgba(123,27,46,.4), rgba(215,25,45,.3)); border-radius:12px; border:2px solid #C9A961; margin:12px 0; }

    /* ============================================================
       OCULTAR BADGES "None" RESIDUALES DE STREAMLIT
       ============================================================ */
    div[data-testid="stMarkdownContainer"] > div:empty,
    div[data-testid="stMarkdownContainer"] > span:empty,
    div[data-testid="stMarkdownContainer"] > p:empty { display:none !important; height:0 !important; margin:0 !important; padding:0 !important; }
    [data-baseweb="badge"]:empty,
    span[data-baseweb="badge"]:not(:has(*)):not([data-testid]) { display:none !important; }
    div.stElementContainer:has(> div > div > span:only-child):not(:has(button)):not(:has(a)):not(:has(img)) { display:none !important; }
    div[data-testid="stMarkdown"] span[style*="color"] { color: transparent !important; }

    @media (max-width: 768px) {
        .stButton > button { min-height:48px !important; font-size:.95rem !important; padding:10px 14px !important; border-radius:12px !important; }
        .stTextInput > div > div > input, .stSelectbox > div > div > div { min-height:44px !important; font-size:16px !important; }
        .block-container { padding-left:12px !important; padding-right:12px !important; padding-top:20px !important; }
        h1 { font-size:1.5rem !important; } h2 { font-size:1.3rem !important; } h3 { font-size:1.1rem !important; }
        .logo-medallon img { width:100px; height:100px; }
        .countdown-num { font-size:5rem; }
        .instruccion-grande { font-size:1.1rem; }
    }
</style>
""", unsafe_allow_html=True)

# ============================================================
# ELIMINADOR DE BADGES "None" RESIDUALES DE STREAMLIT CLOUD
# ============================================================
components.html("""
<script>
(function() {
    const parentDoc = window.parent.document;

    function eliminarNones() {
        try {
            const todos = parentDoc.querySelectorAll('span, div, p, small, b, strong, em, li');
            todos.forEach(function(el) {
                const txt = (el.textContent || '').trim();
                if (txt === 'None' && el.children.length === 0) {
                    el.style.display = 'none';
                    el.style.height = '0';
                    el.style.margin = '0';
                    el.style.padding = '0';
                    el.style.visibility = 'hidden';
                    el.style.opacity = '0';
                    el.setAttribute('aria-hidden', 'true');
                    if (el.parentElement) {
                        el.parentElement.style.display = 'none';
                        el.parentElement.style.height = '0';
                        el.parentElement.style.margin = '0';
                        el.parentElement.style.padding = '0';
                    }
                }
            });

            parentDoc.querySelectorAll('[data-baseweb="badge"], [data-baseweb="tag"]').forEach(function(b) {
                const t = (b.textContent || '').trim();
                if (t === 'None' || t === '' || t === 'null') {
                    b.style.display = 'none';
                    b.style.height = '0';
                    b.setAttribute('aria-hidden', 'true');
                }
            });
        } catch (e) { /* ignorar */ }
    }

    eliminarNones();
    setInterval(eliminarNones, 300);

    try {
        const observer = new MutationObserver(eliminarNones);
        observer.observe(parentDoc.body, { childList: true, subtree: true, characterData: true });
    } catch (e) { /* ignorar */ }
})();
</script>
""", height=0)

st.caption(f"🔧 Versión: {APP_VERSION}")


# ============================================================
# ESCÁNER QR CON AUTOCAPTURA
# ============================================================
if WEBRTC_DISPONIBLE:
    class QRScannerProcessor(VideoProcessorBase):
        def __init__(self):
            self.detector = cv2.QRCodeDetector()
            self._qr = None
            self._lock = threading.Lock()
            self._frame_count = 0
            self._last_frame = None

        def recv(self, frame):
            img = frame.to_ndarray(format="bgr24")
            self._frame_count += 1
            with self._lock:
                self._last_frame = img.copy()
            if self._frame_count % 3 == 0:
                with self._lock:
                    if self._qr is None:
                        try:
                            data, _, _ = self.detector.detectAndDecode(img)
                            if data and len(data.strip()) > 3:
                                self._qr = data.strip()
                        except Exception:
                            pass
            with self._lock:
                detectado = self._qr is not None
            if detectado:
                h, w = img.shape[:2]
                cv2.rectangle(img, (8, 8), (w - 8, h - 8), (0, 255, 0), 6)
                cv2.putText(img, "QR DETECTADO", (25, 55),
                            cv2.FONT_HERSHEY_SIMPLEX, 1.3, (0, 255, 0), 3)
            return av.VideoFrame.from_ndarray(img, format="bgr24")

        def pop_qr(self):
            with self._lock:
                qr = self._qr
                self._qr = None
                return qr

        def capture_frame_jpeg(self):
            with self._lock:
                if self._last_frame is None:
                    return None
                img = self._last_frame.copy()
            try:
                ok, buffer = cv2.imencode(".jpg", img, [cv2.IMWRITE_JPEG_QUALITY, 85])
                if ok:
                    return buffer.tobytes()
            except Exception:
                pass
            return None


def _reset_qr_flow():
    st.session_state["qr_escaneado_actual"] = None
    st.session_state["qr_estado"] = "escaneando"
    st.session_state["qr_tiempo_inicio"] = None
    st.session_state["qr_foto_evidencia"] = None
    st.session_state["qr_vehiculo"] = None


@st.fragment(run_every="1s")
def _poll_qr_scanner():
    processor = st.session_state.get("qr_processor_ref")
    if processor is None:
        return
    if st.session_state.get("qr_estado", "escaneando") != "escaneando":
        return
    qr = processor.pop_qr()
    if qr:
        st.session_state["qr_escaneado_actual"] = qr
        st.session_state["qr_estado"] = "capturando"
        st.session_state["qr_tiempo_inicio"] = time.time()
        st.rerun()


@st.fragment(run_every="1s")
def _fragment_capturando():
    processor = st.session_state.get("qr_processor_ref")
    tiempo_inicio = st.session_state.get("qr_tiempo_inicio") or time.time()
    transcurrido = time.time() - tiempo_inicio
    restante = max(0, SEGUNDOS_AUTOCAPTURA - int(transcurrido))

    if restante > 0:
        st.markdown(
            "<div class='instruccion-grande'>📸 APUNTA LA CÁMARA AL VEHÍCULO AHORA</div>",
            unsafe_allow_html=True
        )
        st.markdown(f"<div class='countdown-num'>{restante}</div>", unsafe_allow_html=True)
        st.progress(min(transcurrido / SEGUNDOS_AUTOCAPTURA, 1.0))
        st.caption("Muévete hacia el vehículo. La foto se tomará automáticamente.")
        return

    foto = None
    if processor is not None:
        foto = processor.capture_frame_jpeg()
    st.session_state["qr_foto_evidencia"] = foto
    st.session_state["qr_estado"] = "confirmando"
    st.rerun()


def _render_escaneando():
    st.markdown("### 📷 Escáner automático de QR")

    st.info(
        "💡 **Instrucciones:**\n\n"
        "1️⃣ Pulsa **START** abajo (permiso de cámara del navegador)\n\n"
        "2️⃣ Apunta la cámara al QR del alumno\n\n"
        "3️⃣ Al detectar, verás una cuenta regresiva de "
        f"**{SEGUNDOS_AUTOCAPTURA} segundos** para mover la cámara hacia el **vehículo**\n\n"
        "4️⃣ La foto se toma automáticamente"
    )

    if WEBRTC_DISPONIBLE:
        st.markdown("🟢 **Cámara lista para activar**")
        try:
            ctx = webrtc_streamer(
                key="qr_scanner_caseta_v2",
                mode=WebRtcMode.SENDRECV,
                video_processor_factory=QRScannerProcessor,
                media_stream_constraints={
                    "video": {
                        "facingMode": "environment",
                        "width": {"ideal": 640},
                        "height": {"ideal": 480},
                    },
                    "audio": False,
                },
                async_processing=True,
                rtc_configuration={
                    "iceServers": [{"urls": ["stun:stun.l.google.com:19302"]}]
                },
            )
            if ctx.video_processor:
                st.session_state["qr_processor_ref"] = ctx.video_processor
            _poll_qr_scanner()
        except Exception as e:
            st.warning(f"⚠️ No se pudo iniciar la cámara: {e}")
    else:
        st.warning("⚠️ El escáner automático no está disponible.")

    st.markdown("---")
    with st.expander("⌨️ Ingresar placas manualmente", expanded=not WEBRTC_DISPONIBLE):
        st.caption("Úsalo si la cámara no funciona.")
        c1, c2 = st.columns([3, 1])
        with c1:
            placas_manual = st.text_input(
                "Placas", key="placas_manual_qr",
                placeholder="Ej. ABC-1234", label_visibility="collapsed"
            ).upper().strip()
        with c2:
            if st.button("🔍 Buscar", use_container_width=True, key="btn_buscar_manual"):
                if placas_manual:
                    st.session_state["qr_escaneado_actual"] = json.dumps({"placas": placas_manual})
                    st.session_state["qr_estado"] = "capturando"
                    st.session_state["qr_tiempo_inicio"] = time.time()
                    st.rerun()
                else:
                    st.error("Ingresa las placas.")


def _render_capturando():
    st.markdown("### ✅ QR detectado")
    _fragment_capturando()


def _render_confirmando(user):
    qr_raw = st.session_state.get("qr_escaneado_actual")
    foto_bytes = st.session_state.get("qr_foto_evidencia")

    try:
        if isinstance(qr_raw, str) and qr_raw.strip().startswith("{"):
            data = json.loads(qr_raw)
        else:
            data = {"placas": str(qr_raw).strip()}
    except Exception:
        data = {"placas": str(qr_raw).strip()}

    placas = (data.get("placas") or "").upper().strip().replace("-", "").replace(" ", "")

    if not placas:
        st.error("❌ QR inválido.")
        if st.button("🔄 Volver a escanear", use_container_width=True, key="qr_reset_1"):
            _reset_qr_flow()
            st.rerun()
        return

    with st.spinner("🔎 Buscando vehículo..."):
        vehiculo = obtener_vehiculo_por_placas(placas)

    if not vehiculo:
        st.error(f"❌ No existe ningún vehículo registrado con las placas **{placas}**.")
        with st.expander("🔍 Ver vehículos registrados (debug)"):
            placas_db = listar_todas_las_placas()
            if not placas_db:
                st.warning("No hay vehículos registrados.")
            else:
                st.caption(f"**{len(placas_db)} vehículo(s):**")
                for p in placas_db:
                    ic = "🚗" if p['tipo'] == 'Auto' else "🏍️"
                    st.markdown(f"- {ic} **{p['placas']}** — {p['nombre_completo']}")
        if st.button("🔄 Volver a escanear", use_container_width=True, key="qr_reset_2"):
            _reset_qr_flow()
            st.rerun()
        return

    st.success("✅ **Datos del alumno**")
    icono = "🚗" if vehiculo['tipo'] == 'Auto' else "🏍️"

    st.markdown(f"### {icono} {vehiculo['nombre_completo']}")
    c1, c2 = st.columns(2)
    with c1:
        st.write(f"**Matrícula:** {vehiculo['matricula'] or 'N/A'}")
        st.write(f"**Carrera:** {vehiculo['carrera'] or 'N/A'}")
    with c2:
        st.write(f"**ID:** {vehiculo['id_estudiante'] or 'N/A'}")
        st.write(f"**Grupo:** {vehiculo['grupo'] or 'N/A'}")

    st.markdown(
        f"<div style='padding:10px 14px; background:linear-gradient(135deg,#7B1B2E,#D7192D); "
        f"border-radius:10px; display:inline-block; margin-bottom:12px;'>"
        f"<b style='color:#C9A961; font-size:1.15rem; letter-spacing:1px;'>"
        f"{vehiculo['tipo']} — {vehiculo['placas']}</b></div>",
        unsafe_allow_html=True
    )

    st.markdown("### 📸 Evidencia capturada")
    if foto_bytes:
        st.image(foto_bytes, caption="Evidencia fotográfica", use_container_width=True)
    else:
        st.warning("⚠️ No se pudo capturar la foto.")

    if st.button("🔄 Recapturar foto", use_container_width=True, key="qr_retry_foto"):
        st.session_state["qr_estado"] = "capturando"
        st.session_state["qr_tiempo_inicio"] = time.time()
        st.session_state["qr_foto_evidencia"] = None
        st.rerun()

    registro_activo = obtener_registro_activo_por_vehiculo(vehiculo['id'])
    accion = "SALIDA" if registro_activo else "ENTRADA"

    if registro_activo:
        try:
            horas_dentro = (datetime.now() - registro_activo['hora_entrada']).total_seconds() / 3600
        except Exception:
            horas_dentro = 0
        if horas_dentro >= 12:
            st.error(f"🚨 Este vehículo lleva **{horas_dentro:.1f}h** dentro.")
        elif horas_dentro >= 8:
            st.warning(f"⚠️ Este vehículo lleva **{horas_dentro:.1f}h** dentro.")
        st.info(f"🟢 Está **DENTRO** desde {registro_activo['hora_entrada']} → Se registrará **SALIDA**")
    else:
        st.info("🔵 **NO** está dentro → Se registrará **ENTRADA**")

    st.markdown(f"### ✅ Confirmar {accion}")

    c_si, c_no = st.columns(2)
    with c_si:
        if st.button(f"✔️ Confirmar {accion}", type="primary", use_container_width=True, key="btn_conf_qr"):
            if not foto_bytes:
                st.error("❌ No hay foto de evidencia.")
            else:
                with st.spinner(f"💾 Registrando {accion.lower()}..."):
                    foto_b64 = imagen_a_base64(foto_bytes)
                    if registro_activo:
                        registrar_salida(registro_activo['id'], vehiculo['tipo'], user['id'], foto_b64)
                        registrar_log(user['id'], "REGISTRAR_SALIDA",
                                      f"Salida de {vehiculo['placas']} ({vehiculo['nombre_completo']}) por QR",
                                      "Super_Registros", registro_activo['id'])
                        msg = f"✅ Salida registrada para {vehiculo['placas']}."
                    else:
                        registrar_entrada(vehiculo['id_usuario'], vehiculo['id'], vehiculo['tipo'],
                                          user['id'], foto_b64)
                        registrar_log(user['id'], "REGISTRAR_ENTRADA",
                                      f"Entrada de {vehiculo['placas']} ({vehiculo['nombre_completo']}) por QR",
                                      "Super_Registros")
                        msg = f"✅ Entrada registrada para {vehiculo['placas']}."
                _reset_qr_flow()
                set_flash("success", msg)
                st.rerun()
    with c_no:
        if st.button("❌ Cancelar", use_container_width=True, key="btn_cancel_qr"):
            _reset_qr_flow()
            st.rerun()


def _seccion_escaner_qr(user):
    estado = st.session_state.get("qr_estado", "escaneando")
    if estado == "escaneando":
        _render_escaneando()
    elif estado == "capturando":
        _render_capturando()
    elif estado == "confirmando":
        _render_confirmando(user)
    else:
        _reset_qr_flow()
        st.rerun()


# ============================================================
# PDF PROFESIONAL
# ============================================================
def generar_pdf_reporte(df, titulo, subtitulo=""):
    class PDF(FPDF):
        def header(self):
            self.set_fill_color(123, 27, 46)
            self.rect(0, 0, 297, 22, "F")
            self.set_y(6); self.set_x(10)
            self.set_font("Helvetica", "B", 12)
            self.set_text_color(201, 169, 97)
            self.cell(0, 6, "CUYPARK - Colegio Universitario de Yahualica", ln=1)
            self.set_font("Helvetica", "", 9)
            self.set_text_color(245, 240, 232)
            self.set_x(10)
            self.cell(0, 5, f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", ln=1)
            self.set_y(26)

        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", "I", 7)
            self.set_text_color(130, 130, 130)
            self.cell(0, 5, f"Página {self.page_no()}", align="C")

    pdf = PDF(orientation="L", unit="mm", format="A4")
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.set_text_color(123, 27, 46)
    pdf.cell(0, 10, titulo, ln=1, align="C")

    if subtitulo:
        pdf.set_font("Helvetica", "I", 10)
        pdf.set_text_color(80, 80, 80)
        pdf.cell(0, 6, subtitulo, ln=1, align="C")
    pdf.ln(3)

    if df.empty:
        pdf.set_font("Helvetica", "", 11)
        pdf.set_text_color(100, 100, 100)
        pdf.cell(0, 10, "No hay registros para mostrar.", ln=1, align="C")
        return bytes(pdf.output())

    page_width = 297 - 20
    cols = list(df.columns)[:9]
    df = df[cols]
    n = len(cols)
    widths = [page_width / n] * n

    pdf.set_fill_color(123, 27, 46)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 8)
    pdf.set_x(10)
    for i, c in enumerate(cols):
        pdf.cell(widths[i], 7, str(c)[:25], border=1, align="C", fill=True)
    pdf.ln()

    pdf.set_font("Helvetica", "", 7)
    pdf.set_text_color(40, 40, 40)
    fill = False
    for _, row in df.iterrows():
        if pdf.get_y() > 190:
            pdf.add_page()
            pdf.set_fill_color(123, 27, 46)
            pdf.set_text_color(255, 255, 255)
            pdf.set_font("Helvetica", "B", 8)
            pdf.set_x(10)
            for i, c in enumerate(cols):
                pdf.cell(widths[i], 7, str(c)[:25], border=1, align="C", fill=True)
            pdf.ln()
            pdf.set_font("Helvetica", "", 7)
            pdf.set_text_color(40, 40, 40)

        if fill:
            pdf.set_fill_color(245, 240, 232)
        else:
            pdf.set_fill_color(255, 255, 255)
        pdf.set_x(10)
        for i, c in enumerate(cols):
            val = row[c]
            if pd.isna(val):
                val = ""
            pdf.cell(widths[i], 6, str(val)[:30], border=1, align="C", fill=True)
        pdf.ln()
        fill = not fill

    return bytes(pdf.output())


# ============================================================
# HELPERS
# ============================================================
def selector_rango_fechas(key_prefix):
    hoy = date.today()
    estado_key = f"{key_prefix}_rango"
    if estado_key not in st.session_state:
        st.session_state[estado_key] = "7d"

    st.caption("**Filtro rápido:**")
    col1, col2, col3, col4 = st.columns(4)
    with col1:
        if st.button("Hoy", key=f"{key_prefix}_btn_hoy", use_container_width=True):
            st.session_state[estado_key] = "hoy"; st.rerun()
    with col2:
        if st.button("7 días", key=f"{key_prefix}_btn_7d", use_container_width=True):
            st.session_state[estado_key] = "7d"; st.rerun()
    with col3:
        if st.button("Este mes", key=f"{key_prefix}_btn_mes", use_container_width=True):
            st.session_state[estado_key] = "mes"; st.rerun()
    with col4:
        if st.button("Todo", key=f"{key_prefix}_btn_todo", use_container_width=True):
            st.session_state[estado_key] = "todo"; st.rerun()

    opcion = st.session_state[estado_key]
    if opcion == "hoy":
        fd, fh = hoy, hoy
    elif opcion == "7d":
        fd, fh = hoy - timedelta(days=7), hoy
    elif opcion == "mes":
        fd, fh = hoy.replace(day=1), hoy
    else:
        fd, fh = None, None

    if opcion == "hoy":
        st.info(f"📅 Mostrando registros de **hoy** ({hoy.strftime('%d/%m/%Y')})")
    elif opcion == "7d":
        st.info(f"📅 **Últimos 7 días** ({fd.strftime('%d/%m/%Y')} → {fh.strftime('%d/%m/%Y')})")
    elif opcion == "mes":
        st.info(f"📅 **{hoy.strftime('%B %Y').capitalize()}**")
    else:
        st.info("📅 **Todos los registros** (sin filtro de fecha)")

    return fd, fh


def mostrar_paginacion_superior(key, por_pagina_options=(10, 20, 50, 100)):
    key_pp = f"{key}_por_pagina"
    if key_pp not in st.session_state:
        st.session_state[key_pp] = 20
    c1, c2 = st.columns([1, 3])
    with c1:
        st.session_state[key_pp] = st.selectbox(
            "Por página", por_pagina_options,
            index=por_pagina_options.index(st.session_state[key_pp]),
            key=f"{key}_pp_sel"
        )
    return st.session_state[key_pp]


def render_paginacion_inferior(key, pagina, total, por_pagina, session_pag_key):
    total_pags = max(1, (total + por_pagina - 1) // por_pagina)
    c1, c2, c3 = st.columns([1, 2, 1])
    with c1:
        if st.button("⬅️ Anterior", key=f"{key}_prev_b", disabled=(pagina <= 1), use_container_width=True):
            st.session_state[session_pag_key] = pagina - 1; st.rerun()
    with c2:
        st.markdown(
            f"<div style='text-align:center; color:#C9A961; padding-top:6px;'>"
            f"Página {pagina} de {total_pags}</div>",
            unsafe_allow_html=True
        )
    with c3:
        if st.button("Siguiente ➡️", key=f"{key}_next_b", disabled=(pagina >= total_pags), use_container_width=True):
            st.session_state[session_pag_key] = pagina + 1; st.rerun()


def notificar_entrada_reciente(user):
    registro = obtener_registro_activo_de_usuario(user['id'])
    if not registro:
        return
    clave = f"toast_entrada_{user['id']}_{registro['id']}"
    if st.session_state.get(clave, False):
        return
    try:
        delta = datetime.now() - registro['hora_entrada']
        minutos = int(delta.total_seconds() / 60)
    except Exception:
        minutos = 999
    ic = "🚗" if registro['tipo'] == 'Auto' else "🏍️"
    try:
        if minutos <= 5:
            st.toast(f"🎉 ¡Bienvenido! Tu {registro['tipo'].lower()} {registro['placas']} acaba de entrar", icon="🎉")
        elif minutos <= 60:
            st.toast(f"{ic} Tu vehículo entró hace {minutos} min", icon="ℹ️")
    except AttributeError:
        pass
    st.session_state[clave] = True


def formatear_tiempo_dentro(hora_entrada):
    try:
        delta = datetime.now() - hora_entrada
        m = int(delta.total_seconds() / 60)
        if m < 1: return "hace unos segundos"
        elif m < 60: return f"hace {m} min"
        else:
            h = m // 60; r = m % 60
            return f"desde hace {h}h" if r == 0 else f"desde hace {h}h {r}min"
    except Exception:
        return ""


def calcular_horas_dentro(hora_entrada):
    try:
        return (datetime.now() - hora_entrada).total_seconds() / 3600
    except Exception:
        return 0


def obtener_badge_alerta(horas):
    if horas >= 12:
        return f'<span class="alert-badge critico">🚨 {horas:.1f}h DENTRO</span>'
    elif horas >= 8:
        return f'<span class="alert-badge advertencia">⚠️ {horas:.1f}h DENTRO</span>'
    return f'<span class="alert-badge ok">✓ {horas:.1f}h</span>'


def validar_nombre(n):
    if not n or not n.strip(): return False, "El nombre es obligatorio."
    if not re.match(r"^[A-Za-zÁÉÍÓÚáéíóúÑñÜü\s\-\.']{2,100}$", n.strip()):
        return False, "Solo letras, espacios, acentos y guiones (2-100 caracteres)."
    return True, ""


def validar_usuario(u):
    if not u or not u.strip(): return False, "El usuario es obligatorio."
    if not re.match(r"^[a-zA-Z0-9._]{3,30}$", u.strip()):
        return False, "3-30 caracteres: letras, números, punto o guión bajo."
    return True, ""


def validar_password(p, minimo=3):
    if not p: return False, "La contraseña es obligatoria."
    if len(p) < minimo: return False, f"Debe tener al menos {minimo} caracteres."
    return True, ""


def validar_telefono(t):
    if not t or not t.strip(): return True, ""
    sd = re.sub(r"\D", "", t)
    if len(sd) != 10: return False, "Debe tener exactamente 10 dígitos."
    return True, ""


def validar_matricula(m):
    if not m or not m.strip(): return True, ""
    if not re.match(r"^[A-Za-z0-9\-]{4,15}$", m.strip()):
        return False, "4-15 caracteres alfanuméricos."
    return True, ""


def validar_id_estudiante(i):
    if not i or not i.strip(): return True, ""
    if not re.match(r"^[A-Za-z0-9\-]{4,20}$", i.strip()):
        return False, "4-20 caracteres alfanuméricos."
    return True, ""


def validar_placas(p):
    if not p or not p.strip(): return False, "Las placas son obligatorias."
    limpio = p.strip().upper().replace("-", "").replace(" ", "")
    if not re.match(r"^[A-Z]{3}[0-9]{3,4}[A-Z]?$", limpio):
        return False, "Formato no válido. Ejemplo: ABC-1234"
    return True, ""


def validar_carrera(c):
    if not c or not c.strip(): return True, ""
    if len(c.strip()) < 2: return False, "Mínimo 2 caracteres."
    if len(c.strip()) > 100: return False, "Máximo 100 caracteres."
    return True, ""


def validar_grupo(g):
    if not g or not g.strip(): return True, ""
    if len(g.strip()) > 20: return False, "Máximo 20 caracteres."
    return True, ""


def mostrar_validacion(v, validador, obligatorio=True):
    if not v or not v.strip():
        return False if obligatorio else True
    ok, msg = validador(v)
    if ok:
        st.markdown('<div class="val-ok">✅ Campo válido</div>', unsafe_allow_html=True)
        return True
    st.markdown(f'<div class="val-error">❌ {msg}</div>', unsafe_allow_html=True)
    return False


def limpiar_campos(keys):
    for k in keys:
        if k in st.session_state:
            del st.session_state[k]


def grafico_dona_ocupacion(ao, al, mo, ml):
    fig = go.Figure(data=[go.Pie(
        labels=["🚗 Autos", "🏍️ Motos", "Libres"],
        values=[ao, mo, al + ml], hole=0.6,
        marker=dict(colors=[COLOR_VINO_CLARO, COLOR_DORADO, "#2a2a2a"], line=dict(color="#0d0407", width=2)),
        textinfo="label+percent", textfont=dict(size=11, color=COLOR_CREMA),
        hovertemplate="<b>%{label}</b><br>Cantidad: %{value}<br>%{percent}<extra></extra>"
    )])
    fig.update_layout(
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
        font=dict(color=COLOR_CREMA, size=11),
        margin=dict(l=10, r=10, t=10, b=10), height=280,
        legend=dict(orientation="h", yanchor="bottom", y=-0.1, xanchor="center", x=0.5, font=dict(size=10, color=COLOR_CREMA)),
        annotations=[dict(text=f"<b>{ao + mo}</b><br><span style='font-size:10px'>Ocupados</span>",
                          x=0.5, y=0.5, font=dict(size=18, color=COLOR_DORADO), showarrow=False)]
    )
    return fig


def grafico_barras_horas(df):
    fig = go.Figure(data=[go.Bar(
        x=df["hora_str"], y=df["entradas"],
        marker=dict(color=df["entradas"], colorscale=[[0, COLOR_VINO], [1, COLOR_DORADO]], line=dict(color=COLOR_DORADO, width=1)),
        text=df["entradas"], textposition="outside", textfont=dict(color=COLOR_CREMA, size=10),
        hovertemplate="<b>%{x}</b><br>Entradas: %{y}<extra></extra>"
    )])
    fig.update_layout(
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
        font=dict(color=COLOR_CREMA, size=11),
        xaxis=dict(gridcolor="rgba(201,169,97,0.1)", tickfont=dict(color=COLOR_CREMA, size=10)),
        yaxis=dict(gridcolor="rgba(201,169,97,0.1)", tickfont=dict(color=COLOR_CREMA, size=10)),
        margin=dict(l=10, r=10, t=20, b=10), height=280, showlegend=False
    )
    return fig


def grafico_linea_tendencia(df):
    fig = go.Figure()
    fig.add_trace(go.Scatter(
        x=df["fecha_str"], y=df["entradas"], mode="lines+markers",
        line=dict(color=COLOR_DORADO, width=3, shape="spline"),
        marker=dict(color=COLOR_VINO_CLARO, size=10, line=dict(color=COLOR_DORADO, width=2)),
        fill="tozeroy", fillcolor="rgba(123,27,46,0.25)",
        hovertemplate="<b>%{x}</b><br>Entradas: %{y}<extra></extra>"
    ))
    fig.update_layout(
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
        font=dict(color=COLOR_CREMA, size=11),
        xaxis=dict(gridcolor="rgba(201,169,97,0.1)", tickfont=dict(color=COLOR_CREMA, size=10)),
        yaxis=dict(gridcolor="rgba(201,169,97,0.1)", tickfont=dict(color=COLOR_CREMA, size=10)),
        margin=dict(l=10, r=10, t=20, b=10), height=280, showlegend=False
    )
    return fig


def grafico_barras_carreras(df):
    fig = go.Figure(data=[go.Bar(
        y=df["carrera"], x=df["visitas"], orientation="h",
        marker=dict(color=df["visitas"], colorscale=[[0, COLOR_VINO], [1, COLOR_DORADO]], line=dict(color=COLOR_DORADO, width=1)),
        text=df["visitas"], textposition="outside", textfont=dict(color=COLOR_CREMA, size=10),
        hovertemplate="<b>%{y}</b><br>Visitas: %{x}<extra></extra>"
    )])
    fig.update_layout(
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
        font=dict(color=COLOR_CREMA, size=11),
        xaxis=dict(gridcolor="rgba(201,169,97,0.1)", tickfont=dict(color=COLOR_CREMA, size=10)),
        yaxis=dict(gridcolor="rgba(201,169,97,0.1)", tickfont=dict(color=COLOR_CREMA, size=10)),
        margin=dict(l=10, r=10, t=20, b=10), height=280, showlegend=False
    )
    return fig


def generar_pdf_qr(user, vehiculo, qr_bytes):
    class PDF(FPDF):
        def header(self):
            self.set_fill_color(123, 27, 46)
            self.rect(0, 0, 210, 28, "F")
            try:
                with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as _f:
                    _f.write(urllib.request.urlopen(LOGO_ESCUDO_URL, timeout=5).read())
                    lp = _f.name
                self.image(lp, x=10, y=4, w=20)
                try: os.unlink(lp)
                except Exception: pass
            except Exception:
                pass
            self.set_y(6); self.set_x(35)
            self.set_font("Helvetica", "B", 10)
            self.set_text_color(201, 169, 97)
            self.cell(0, 5, "COLEGIO UNIVERSITARIO", ln=1, align="L")
            self.set_x(35)
            self.set_font("Helvetica", "B", 14)
            self.cell(0, 6, "DE YAHUALICA", ln=1, align="L")
            self.set_y(30)

        def footer(self):
            self.set_y(-18)
            self.set_font("Helvetica", "I", 7)
            self.set_text_color(130, 130, 130)
            self.cell(0, 5, "CUYPARK - Sistema de Estacionamiento Inteligente", align="C", ln=1)
            self.cell(0, 5, f"Documento generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')}", align="C")

    pdf = PDF(orientation="P", unit="mm", format="A4")
    pdf.add_page()
    pdf.set_y(35)
    pdf.set_font("Helvetica", "B", 22)
    pdf.set_text_color(123, 27, 46)
    pdf.cell(0, 10, "CUYPARK", ln=1, align="C")
    pdf.set_font("Helvetica", "", 11)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(0, 6, "Boleto Digital de Estacionamiento", ln=1, align="C")
    pdf.ln(4)
    pdf.set_draw_color(201, 169, 97); pdf.set_line_width(0.5)
    pdf.line(60, pdf.get_y(), 150, pdf.get_y())
    pdf.ln(6)

    pdf.set_font("Helvetica", "B", 12); pdf.set_text_color(123, 27, 46)
    pdf.cell(0, 7, "Datos del Alumno", ln=1)
    pdf.set_font("Helvetica", "", 10); pdf.set_text_color(40, 40, 40)
    for et, v in [("Nombre completo:", user.get('nombre_completo') or 'N/A'),
                  ("ID Estudiante:", user.get('id_estudiante') or 'N/A'),
                  ("Matrícula:", user.get('matricula') or 'N/A'),
                  ("Carrera:", user.get('carrera') or 'N/A'),
                  ("Grupo:", user.get('grupo') or 'N/A')]:
        pdf.cell(45, 6, et, border=0)
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(0, 6, str(v), ln=1)
        pdf.set_font("Helvetica", "", 10)

    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 12); pdf.set_text_color(123, 27, 46)
    pdf.cell(0, 7, "Datos del Vehículo", ln=1)
    pdf.set_font("Helvetica", "", 10); pdf.set_text_color(40, 40, 40)
    for et, v, ep in [("Tipo:", vehiculo.get('tipo') or 'N/A', False),
                       ("Placas:", vehiculo.get('placas') or 'N/A', True),
                       ("Marca:", vehiculo.get('marca') or 'N/A', False),
                       ("Modelo:", vehiculo.get('modelo') or 'N/A', False),
                       ("Color:", vehiculo.get('color') or 'N/A', False)]:
        pdf.cell(45, 6, et, border=0)
        pdf.set_font("Helvetica", "B", 12 if ep else 10)
        if ep: pdf.set_text_color(123, 27, 46)
        pdf.cell(0, 6, str(v), ln=1)
        pdf.set_font("Helvetica", "", 10); pdf.set_text_color(40, 40, 40)

    pdf.ln(8)
    pdf.set_font("Helvetica", "B", 12); pdf.set_text_color(123, 27, 46)
    pdf.cell(0, 7, "Código QR de Acceso", ln=1, align="C")

    with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmp:
        tmp.write(qr_bytes); tp = tmp.name
    qs = 90
    qx = (210 - qs) / 2
    qy = pdf.get_y() + 3
    pdf.set_draw_color(201, 169, 97); pdf.set_line_width(0.6)
    pdf.rect(qx - 3, qy - 3, qs + 6, qs + 6)
    pdf.image(tp, x=qx, y=qy, w=qs, h=qs)
    try: os.unlink(tp)
    except Exception: pass
    pdf.set_y(qy + qs + 8)
    pdf.set_font("Helvetica", "I", 9); pdf.set_text_color(100, 100, 100)
    pdf.multi_cell(0, 5,
        "Presenta este código QR en la caseta al entrar y salir. "
        "Puedes imprimirlo o mostrarlo desde tu móvil.", align="C")
    return bytes(pdf.output())


# ============================================================
# SESIÓN
# ============================================================
if 'usuario' not in st.session_state: st.session_state.usuario = None
if 'qr_generado' not in st.session_state: st.session_state.qr_generado = None
if 'flash' not in st.session_state: st.session_state.flash = None
if 'qr_escaneado_actual' not in st.session_state: st.session_state.qr_escaneado_actual = None
if 'qr_processor_ref' not in st.session_state: st.session_state.qr_processor_ref = None
if 'qr_estado' not in st.session_state: st.session_state.qr_estado = "escaneando"
if 'qr_tiempo_inicio' not in st.session_state: st.session_state.qr_tiempo_inicio = None
if 'qr_foto_evidencia' not in st.session_state: st.session_state.qr_foto_evidencia = None
if 'qr_vehiculo' not in st.session_state: st.session_state.qr_vehiculo = None

if 'tabla_mensajes_ok' not in st.session_state:
    try:
        crear_tabla_mensajes()
        st.session_state['tabla_mensajes_ok'] = True
    except Exception:
        st.session_state['tabla_mensajes_ok'] = False


def mostrar_flash():
    if st.session_state.flash:
        t, x = st.session_state.flash
        if t == "success": st.success(x)
        elif t == "error": st.error(x)
        elif t == "warning": st.warning(x)
        elif t == "info": st.info(x)
        st.session_state.flash = None


def set_flash(t, x): st.session_state.flash = (t, x)


def mostrar_branding_login():
    c1, c2, c3 = st.columns([1, 2, 1])
    with c2:
        st.markdown(f"""
            <div style="text-align: center; margin-bottom: 24px;">
                <div class="logo-medallon"><img src="{LOGO_ESCUDO_URL}" alt="CUY"></div>
                <div style="color: #C9A961; font-size: 0.7rem; letter-spacing: 4px; text-transform: uppercase; font-weight: 600; margin-top: 4px;">Colegio Universitario</div>
                <div style="color: #C9A961; font-size: 1.6rem; font-weight: 800; letter-spacing: 3px; margin-top: 2px;">DE YAHUALICA</div>
                <div style="width: 80px; height: 2px; background: linear-gradient(90deg, transparent, #C9A961, transparent); margin: 12px auto 0 auto;"></div>
            </div>
        """, unsafe_allow_html=True)


def mostrar_marca_cudy():
    return f"""
        <div class="brand-cudy">
            <div class="texto">
                <div class="linea1">Colegio Universitario</div>
                <div class="linea2">DE YAHUALICA</div>
            </div>
            <div class="logo-barra" style="margin-right: 0;">
                <img src="{LOGO_ESCUDO_URL}" alt="CUY">
            </div>
        </div>
    """


def cb_mostrar_confirm(k): st.session_state[k] = True
def cb_ocultar_confirm(k): st.session_state[k] = False


def cb_eliminar_vehiculo(vid, uid, tipo, placas):
    ok, msg = eliminar_vehiculo(vid)
    st.session_state[f"confirmar_elim_veh_{vid}"] = False
    if ok:
        registrar_log(uid, "ELIMINAR_VEHICULO", f"Vehículo {tipo} {placas} eliminado", "Super_Vehiculos", vid)
        st.session_state.flash = ("success", "Vehículo eliminado.")
    else:
        st.session_state.flash = ("error", msg)


def cb_activar_usuario(uid, aid, au, tu, tn):
    ok, msg = activar_usuario(uid)
    if ok:
        registrar_log(aid, "REACTIVAR_USUARIO", f"Cuenta @{tu} reactivada por @{au}", "Super_Usuarios", uid)
        st.session_state.flash = ("success", msg)
    else:
        st.session_state.flash = ("error", msg)


def cb_desactivar_usuario(uid, aid, au, tu, tn):
    ok, msg = desactivar_usuario(uid)
    st.session_state[f"confirmar_desactivar_{uid}"] = False
    if ok:
        registrar_log(aid, "DESACTIVAR_USUARIO", f"Cuenta @{tu} desactivada por @{au}", "Super_Usuarios", uid)
        st.session_state.flash = ("success", msg)
    else:
        st.session_state.flash = ("error", msg)


def cb_eliminar_usuario(uid, aid, au, tu, tn):
    ok, msg = eliminar_usuario(uid)
    st.session_state[f"confirmar_eliminar_{uid}"] = False
    if ok:
        registrar_log(aid, "ELIMINAR_USUARIO", f"Cuenta @{tu} eliminada por @{au}", "Super_Usuarios", uid)
        st.session_state.flash = ("success", msg)
    else:
        st.session_state.flash = ("error", msg)


def cerrar_sesion():
    st.session_state.usuario = None
    st.session_state.qr_generado = None
    st.session_state.flash = None
    _reset_qr_flow()
    st.rerun()


def mostrar_mi_cuenta(user):
    with st.expander("🔧 Mi cuenta — Contraseña y contacto"):
        st.markdown("#### 📱 Actualizar mi teléfono")
        st.caption("Solo el teléfono puede ser editado por ti.")
        ta = user.get('telefono') or ""
        st.caption(f"Teléfono actual: **{ta if ta else '(sin teléfono)'}**")
        nt = st.text_input("Nuevo teléfono (10 dígitos)", value=ta,
                           key=f"mt_tel_{user['id']}", placeholder="444-123-4567")
        tok = False
        if nt and nt.strip():
            ok, msg = validar_telefono(nt)
            if ok:
                st.markdown('<div class="val-ok">✅ Teléfono válido</div>', unsafe_allow_html=True)
                tok = True
            else:
                st.markdown(f'<div class="val-error">❌ {msg}</div>', unsafe_allow_html=True)
        else:
            st.markdown('<div class="val-error">❌ El teléfono es obligatorio</div>', unsafe_allow_html=True)
        tm = (nt.strip() != ta.strip()) if nt else False
        ttok = tok and tm
        if st.button("💾 Guardar teléfono", use_container_width=True, type="primary",
                     disabled=not ttok, key=f"mt_btn_{user['id']}"):
            ok, msg = actualizar_telefono_usuario(user['id'], nt.strip())
            if ok:
                registrar_log(user['id'], "ACTUALIZAR_TELEFONO",
                              f"@{user['usuario']} actualizó teléfono", "Super_Usuarios", user['id'])
                st.session_state.usuario['telefono'] = nt.strip()
                limpiar_campos([f"mt_tel_{user['id']}"])
                set_flash("success", f"✅ {msg}")
                st.rerun()
            else:
                st.error(f"❌ {msg}")

        st.markdown("---")
        st.markdown("#### 🔑 Cambiar contraseña")
        pa = st.text_input("Contraseña actual", type="password", key=f"cp_act_{user['id']}")
        pn = st.text_input("Nueva contraseña", type="password", key=f"cp_new_{user['id']}")
        pn_ok = False
        if pn:
            if len(pn) < 3:
                st.markdown('<div class="val-error">❌ Mínimo 3 caracteres</div>', unsafe_allow_html=True)
            else:
                st.markdown('<div class="val-ok">✅ Válida</div>', unsafe_allow_html=True)
                pn_ok = True
        pc = st.text_input("Confirmar nueva", type="password", key=f"cp_conf_{user['id']}")
        pc_ok = False
        if pc:
            if pc == pn:
                st.markdown('<div class="val-ok">✅ Coinciden</div>', unsafe_allow_html=True)
                pc_ok = True
            else:
                st.markdown('<div class="val-error">❌ No coinciden</div>', unsafe_allow_html=True)
        tpk = bool(pa) and pn_ok and pc_ok
        if st.button("🔐 Cambiar contraseña", use_container_width=True, type="primary",
                     disabled=not tpk, key=f"cp_btn_{user['id']}"):
            ok, msg = cambiar_password_usuario(user['id'], pa, pn)
            if ok:
                registrar_log(user['id'], "CAMBIAR_PASSWORD",
                              f"@{user['usuario']} cambió contraseña", "Super_Usuarios", user['id'])
                limpiar_campos([f"cp_act_{user['id']}", f"cp_new_{user['id']}", f"cp_conf_{user['id']}"])
                set_flash("success", f"✅ {msg}")
                st.rerun()
            else:
                st.error(f"❌ {msg}")


def pantalla_login():
    mostrar_branding_login()
    st.markdown('<p class="titulo-principal">CUYPARK</p>', unsafe_allow_html=True)
    st.markdown('<p class="subtitulo">Sistema de Estacionamiento Inteligente — CUY</p>', unsafe_allow_html=True)
    mostrar_flash()

    with st.spinner("🔐 Verificando sistema..."):
        ha = contar_admins() > 0

    if not ha:
        with st.expander("🚨 Configuración inicial: Crear el primer Administrador", expanded=True):
            st.warning("No existe ningún administrador.")
            ua = st.text_input("Usuario admin", key="pa_u")
            u_ok = mostrar_validacion(ua, validar_usuario, obligatorio=True)
            pas = st.text_input("Contraseña", type="password", key="pa_p")
            p_ok = mostrar_validacion(pas, validar_password, obligatorio=True)
            na = st.text_input("Nombre completo", key="pa_n")
            n_ok = mostrar_validacion(na, validar_nombre, obligatorio=True)
            tok = u_ok and p_ok and n_ok
            if st.button("Crear Administrador", use_container_width=True, type="primary",
                         disabled=not tok, key="pa_btn"):
                if usuario_existe(ua.lower().strip()):
                    st.error("❌ Ese usuario ya existe")
                else:
                    crear_usuario(usuario=ua.lower().strip(), password=pas,
                                  rol='admin', tipo_usuario='administrativo',
                                  nombre_completo=na.strip())
                    limpiar_campos(['pa_u', 'pa_p', 'pa_n'])
                    set_flash("success", "✅ Administrador creado.")
                    st.rerun()

    st.markdown("### 🔑 Iniciar Sesión")
    usuario = st.text_input("Usuario", key="login_u")
    password = st.text_input("Contraseña", type="password", key="login_p")

    if st.button("Ingresar al Sistema", use_container_width=True, type="primary", key="login_btn"):
        if not usuario or not password:
            st.error("Completa todos los campos")
        else:
            with st.spinner("🔓 Autenticando..."):
                user, error = autenticar(usuario.lower().strip(), password)
            if user:
                st.session_state.usuario = user
                registrar_log(user['id'], "INICIO_SESION",
                              f"@{user['usuario']} inició sesión", "Super_Usuarios", user['id'])
                st.rerun()
            else:
                if "🔒" in error:
                    st.error(f"{error}")
                    st.info("💡 Contacta al administrador.")
                else:
                    st.warning(f"⚠️ {error}")

    st.caption("🔒 Las cuentas son creadas por el administrador.")


@st.fragment(run_every="15s")
def _contadores_alumno():
    esp = obtener_espacios()
    a = next((e for e in esp if e['tipo'] == 'Auto'), None)
    m = next((e for e in esp if e['tipo'] == 'Moto'), None)
    c1, c2 = st.columns(2)
    with c1:
        if a:
            d = a['capacidad_total'] - a['ocupados']
            st.markdown(f'<div class="contador-card"><h4>🚗 Autos</h4><div class="numero">{d} de {a["capacidad_total"]}</div></div>', unsafe_allow_html=True)
    with c2:
        if m:
            d = m['capacidad_total'] - m['ocupados']
            st.markdown(f'<div class="contador-card"><h4>🏍️ Motos</h4><div class="numero">{d} de {m["capacidad_total"]}</div></div>', unsafe_allow_html=True)


@st.fragment(run_every="15s")
def _contadores_caseta():
    esp = obtener_espacios()
    a = next((e for e in esp if e['tipo'] == 'Auto'), None)
    m = next((e for e in esp if e['tipo'] == 'Moto'), None)
    c1, c2 = st.columns(2)
    with c1:
        if a: st.metric("🚗 Autos dentro", f"{a['ocupados']} / {a['capacidad_total']}")
    with c2:
        if m: st.metric("🏍️ Motos dentro", f"{m['ocupados']} / {m['capacidad_total']}")


@st.fragment(run_every="5m")
def _verificar_alertas_programadas(user):
    try:
        al = obtener_vehiculos_alerta(12)
        if al:
            k = f"alerta_{date.today().isoformat()}_{len(al)}"
            if not st.session_state.get(k):
                st.toast(f"🚨 {len(al)} vehículo(s) llevan +12h dentro.", icon="🚨")
                st.session_state[k] = True
    except Exception:
        pass


@st.fragment(run_every="15s")
def _dashboard_datos_vivo():
    esp = obtener_espacios()
    autos = next((e for e in esp if e['tipo'] == 'Auto'), None)
    motos = next((e for e in esp if e['tipo'] == 'Moto'), None)

    cl, ct = st.columns([1, 6])
    with cl:
        st.markdown(f"""
            <div style="background: radial-gradient(circle, #FFFFFF 0%, #F5F0E8 100%); border-radius: 50%; padding: 4px;
                border: 3px solid {COLOR_DORADO}; box-shadow: 0 0 12px rgba(201,169,97,.5);
                width:70px; height:70px; display:flex; align-items:center; justify-content:center;">
                <img src="{LOGO_ESCUDO_URL}" style="width:60px; height:60px; border-radius:50%;" alt="CUY">
            </div>
        """, unsafe_allow_html=True)
    with ct:
        st.markdown(f"""
            <div style="padding-top: 8px;">
                <div style="color:{COLOR_DORADO}; font-size:1.6rem; font-weight:800; letter-spacing:1px;">📊 Panel de Control</div>
                <div style="color:{COLOR_CREMA}; font-size:.85rem; opacity:.7; margin-top:-4px;">Colegio Universitario de Yahualica · Actualizado cada 15s</div>
            </div>
        """, unsafe_allow_html=True)

    st.markdown("---")
    c1, c2 = st.columns(2)
    with c1:
        if autos:
            l = autos['capacidad_total'] - autos['ocupados']
            st.markdown(f'<div class="kpi-card"><div class="kpi-label">🚗 Autos dentro</div><div class="kpi-value">{autos["ocupados"]}</div><div class="kpi-delta delta-neutral">{l} lugares libres</div></div>', unsafe_allow_html=True)
    with c2:
        if motos:
            l = motos['capacidad_total'] - motos['ocupados']
            st.markdown(f'<div class="kpi-card"><div class="kpi-label">🏍️ Motos dentro</div><div class="kpi-value">{motos["ocupados"]}</div><div class="kpi-delta delta-neutral">{l} lugares libres</div></div>', unsafe_allow_html=True)

    st.markdown("#### 📊 Ocupación actual")
    if autos and autos['capacidad_total'] > 0:
        p = (autos['ocupados'] / autos['capacidad_total']) * 100
        st.markdown(f'<div class="ocupacion-label"><span>🚗 Autos</span><b>{autos["ocupados"]}/{autos["capacidad_total"]} ({p:.1f}%)</b></div>', unsafe_allow_html=True)
        st.progress(min(p / 100, 1.0))
    if motos and motos['capacidad_total'] > 0:
        p = (motos['ocupados'] / motos['capacidad_total']) * 100
        st.markdown(f'<div class="ocupacion-label"><span>🏍️ Motos</span><b>{motos["ocupados"]}/{motos["capacidad_total"]} ({p:.1f}%)</b></div>', unsafe_allow_html=True)
        st.progress(min(p / 100, 1.0))

    st.markdown("---")
    registros = obtener_todos_los_registros()

    if registros:
        df = pd.DataFrame(registros)
        df['hora_entrada'] = pd.to_datetime(df['hora_entrada'], errors='coerce')
        df['fecha'] = df['hora_entrada'].dt.date
        hoy = date.today(); ayer = hoy - timedelta(days=1)

        st.markdown("#### 📈 Actividad reciente")
        hc = len(df[df['fecha'] == hoy]); ac = len(df[df['fecha'] == ayer])
        c1, c2, c3 = st.columns(3)
        with c1:
            d = hc - ac
            dt = f"{d:+d} ({((d/ac)*100):+.0f}%)" if ac > 0 else "Sin datos"
            st.metric("Entradas hoy", hc, dt)
        with c2: st.metric("Entradas ayer", ac)
        with c3: st.metric("Usuarios totales", len(obtener_todos_usuarios()))

        st.markdown("---")
        if autos and motos:
            st.markdown("#### 🍩 Distribución de ocupación")
            al = autos['capacidad_total'] - autos['ocupados']
            ml = motos['capacidad_total'] - motos['ocupados']
            st.plotly_chart(grafico_dona_ocupacion(autos['ocupados'], al, motos['ocupados'], ml),
                            use_container_width=True, config={"displayModeBar": False})

        st.markdown("#### ⏰ Horas de mayor demanda (7 días)")
        h7 = hoy - timedelta(days=6); df7 = df[df['fecha'] >= h7].copy()
        if not df7.empty:
            df7['hora'] = df7['hora_entrada'].dt.hour
            hp = df7.groupby('hora').size().reset_index(name='entradas')
            th = pd.DataFrame({'hora': range(6, 22)})
            hp = th.merge(hp, on='hora', how='left').fillna(0)
            hp['hora_str'] = hp['hora'].apply(lambda h: f"{int(h):02d}:00")
            st.plotly_chart(grafico_barras_horas(hp), use_container_width=True, config={"displayModeBar": False})
        else:
            st.caption("Sin datos.")

        st.markdown("#### 📅 Tendencia 30 días")
        h30 = hoy - timedelta(days=29); df30 = df[df['fecha'] >= h30].copy()
        if not df30.empty:
            ed = df30.groupby('fecha').size().reset_index(name='entradas')
            tf = pd.DataFrame({'fecha': pd.date_range(h30, hoy).date})
            ed = tf.merge(ed, on='fecha', how='left').fillna(0)
            ed['fecha_str'] = pd.to_datetime(ed['fecha']).dt.strftime('%d/%m')
            st.plotly_chart(grafico_linea_tendencia(ed), use_container_width=True, config={"displayModeBar": False})
        else:
            st.caption("Sin datos.")

        st.markdown("#### 🎓 Top 5 carreras")
        pc = df[df['carrera'].notna()].groupby('carrera').size().reset_index(name='visitas')
        pc = pc.sort_values('visitas', ascending=True).tail(5)
        if not pc.empty:
            st.plotly_chart(grafico_barras_carreras(pc), use_container_width=True, config={"displayModeBar": False})
        else:
            st.caption("Sin datos.")
    else:
        st.info("Aún no hay registros.")

    st.markdown("---")
    st.markdown("### 🚘 Vehículos dentro ahora")
    dentro_list = obtener_vehiculos_dentro()
    if not dentro_list:
        st.info("No hay vehículos dentro.")
    else:
        crit = [v for v in dentro_list if calcular_horas_dentro(v['hora_entrada']) >= 12]
        adv = [v for v in dentro_list if 8 <= calcular_horas_dentro(v['hora_entrada']) < 12]
        if crit:
            st.markdown(f'<div class="alert-card"><div style="color:#ff4444; font-weight:800; font-size:1rem;">🚨 {len(crit)} vehículo(s) con +12h dentro</div></div>', unsafe_allow_html=True)
        if adv:
            with st.expander(f"⚠️ {len(adv)} vehículo(s) con 8-12h"):
                for v in adv:
                    h = calcular_horas_dentro(v['hora_entrada'])
                    i = "🚗" if v['tipo'] == 'Auto' else "🏍️"
                    with st.container(border=True):
                        ca, cb = st.columns([3, 2])
                        with ca:
                            st.markdown(f"**{i} {v['placas']}** — {v['nombre_completo']}")
                            st.caption(f"Entrada: {v['hora_entrada']}")
                        with cb:
                            st.markdown(obtener_badge_alerta(h), unsafe_allow_html=True)
        st.markdown(f"#### Todos ({len(dentro_list)})")
        for v in dentro_list:
            h = calcular_horas_dentro(v['hora_entrada'])
            i = "🚗" if v['tipo'] == 'Auto' else "🏍️"
            with st.container(border=True):
                ca, cb = st.columns([3, 2])
                with ca:
                    st.markdown(f"**{i} {v['placas']}** — {v['nombre_completo']}")
                    st.caption(f"Matrícula: {v['matricula'] or 'N/A'} | Entrada: {v['hora_entrada']}")
                with cb:
                    st.markdown(obtener_badge_alerta(h), unsafe_allow_html=True)


def _render_lista_vehiculos_dentro(user):
    buscar = st.text_input("🔍 Buscar vehículo", key="buscar_dentro_caseta",
                           placeholder="Ej. ABC-1234 o Juan Pérez")
    dentro = obtener_vehiculos_dentro()
    if not dentro:
        st.info("No hay vehículos dentro.")
        return

    if buscar and buscar.strip():
        t = buscar.strip().lower()
        df_ = [v for v in dentro if
               t in (v['placas'] or '').lower() or
               t in (v['nombre_completo'] or '').lower() or
               t in (v['matricula'] or '').lower()]
        if len(df_) < len(dentro):
            st.caption(f"🔎 **{len(df_)}** de **{len(dentro)}**")
        dentro = df_
    else:
        cr = sum(1 for v in dentro if calcular_horas_dentro(v['hora_entrada']) >= 12)
        ad = sum(1 for v in dentro if 8 <= calcular_horas_dentro(v['hora_entrada']) < 12)
        if cr or ad:
            c1, c2 = st.columns(2)
            with c1:
                if cr: st.markdown(f'<div class="alert-badge critico">🚨 {cr} con +12h</div>', unsafe_allow_html=True)
            with c2:
                if ad: st.markdown(f'<div class="alert-badge advertencia">⚠️ {ad} con +8h</div>', unsafe_allow_html=True)
        st.write(f"**Total: {len(dentro)}**")

    if not dentro:
        st.info("No hay vehículos que coincidan.")
        return

    for v in dentro:
        h = calcular_horas_dentro(v['hora_entrada'])
        with st.container(border=True):
            i = "🚗" if v['tipo'] == 'Auto' else "🏍️"
            ci, cb = st.columns([3, 2])
            with ci:
                st.markdown(f"**{i} {v['placas']}** — {v['nombre_completo']}")
                st.caption(f"Matrícula: {v['matricula'] or 'N/A'} | Entrada: {v['hora_entrada']}")
            with cb:
                st.markdown(obtener_badge_alerta(h), unsafe_allow_html=True)
            if h >= 12:
                st.warning(f"⚠️ Este vehículo lleva **{h:.1f} horas** estacionado.")
            if st.button("🚪 Registrar Salida", key=f"sal_{v['id_registro']}", use_container_width=True):
                st.session_state[f"salida_rapida_{v['id_registro']}"] = True
                st.rerun()

        if st.session_state.get(f"salida_rapida_{v['id_registro']}", False):
            st.markdown("#### 📸 Evidencia de salida")
            foto_sal = st.camera_input(f"Evidencia — {v['placas']}", key=f"cam_sal_{v['id_registro']}")
            c1, c2 = st.columns(2)
            with c1:
                if st.button("✅ Confirmar", key=f"conf_sal_{v['id_registro']}", type="primary", use_container_width=True):
                    if not foto_sal:
                        st.error("Toma la foto de evidencia.")
                    else:
                        with st.spinner("💾 Guardando salida..."):
                            registrar_salida(v['id_registro'], v['tipo'], user['id'],
                                             imagen_a_base64(foto_sal.getvalue()))
                            registrar_log(user['id'], "REGISTRAR_SALIDA",
                                          f"Salida de {v['placas']} ({v['nombre_completo']})",
                                          "Super_Registros", v['id_registro'])
                        st.session_state[f"salida_rapida_{v['id_registro']}"] = False
                        set_flash("success", f"✅ Salida registrada para {v['placas']}.")
                        st.rerun()
            with c2:
                if st.button("❌ Cancelar", key=f"canc_sal_{v['id_registro']}", use_container_width=True):
                    st.session_state[f"salida_rapida_{v['id_registro']}"] = False
                    st.rerun()


@st.fragment(run_every="10s")
def _render_dentro_auto(user):
    _render_lista_vehiculos_dentro(user)


@st.fragment
def _render_dentro_manual(user):
    _render_lista_vehiculos_dentro(user)


# ============================================================
# PANEL ALUMNO
# ============================================================
def panel_alumno():
    user = st.session_state.usuario
    st.markdown(f'<div class="panel-header">🎓 Alumno — {user["nombre_completo"]}</div>', unsafe_allow_html=True)
    mostrar_flash()
    notificar_entrada_reciente(user)

    nl = contar_mensajes_no_leidos(user['id'])
    if nl > 0:
        st.info(f"💬 Tienes **{nl}** mensaje(s) sin leer.")

    with st.expander("👤 Ver mi perfil"):
        st.write(f"**Usuario:** {user['usuario']}")
        st.write(f"**ID Estudiante:** {user['id_estudiante'] or 'N/A'}")
        st.write(f"**Matrícula:** {user['matricula'] or 'N/A'}")
        st.write(f"**Carrera:** {user['carrera'] or 'N/A'}")
        st.write(f"**Grupo:** {user['grupo'] or 'N/A'}")
        st.write(f"**Teléfono:** {user['telefono'] or 'N/A'}")

    mostrar_mi_cuenta(user)

    tm = f"💬 Mensajes ({nl} sin leer)" if nl > 0 else "💬 Mensajes"
    with st.expander(tm):
        msj = obtener_mensajes_para_usuario(user['id'])
        if not msj:
            st.info("No tienes mensajes.")
        else:
            for m in msj:
                il = "📬" if not m['leido'] else "📭"
                ti = {"mensaje": "💬", "alerta": "🚨", "aviso": "📢"}.get(m['tipo'], "💬")
                with st.container(border=True):
                    st.markdown(f"**{il} {ti} {m['asunto']}**")
                    st.caption(f"De: {m['remitente_nombre']} ({m['remitente_rol']}) · {m['fecha'].strftime('%d/%m/%Y %H:%M')}")
                    if m.get('es_broadcast'): st.caption("📢 General")
                    with st.expander("Ver mensaje"):
                        st.write(m['cuerpo'])
                        if not m['leido']:
                            if st.button("✅ Marcar leído", key=f"al_leido_{m['id']}", use_container_width=True):
                                marcar_mensaje_leido(m['id'])
                                st.rerun()

    st.markdown("### 🅿️ Lugares disponibles")
    st.caption("🟢 Actualizándose cada 15s")
    _contadores_alumno()
    st.markdown("---")

    ra = obtener_registro_activo_de_usuario(user['id'])
    if ra:
        try:
            dm = datetime.now() - ra['hora_entrada']
            mins = int(dm.total_seconds() / 60)
        except Exception:
            mins = 999
        if mins <= 5:
            st.markdown(f"""
                <div class="bienvenida-card">
                    <h3>🎉 ¡Bienvenido al estacionamiento!</h3>
                    <p>Tu <b>{ra['tipo'].lower()}</b> acaba de entrar:</p>
                    <p class="placas">{ra['placas']}</p>
                </div>
            """, unsafe_allow_html=True)
        else:
            tt = formatear_tiempo_dentro(ra['hora_entrada'])
            st.success(f"✅ Vehículo **DENTRO**: {ra['placas']} ({ra['tipo']}) — {tt}")
            st.caption(f"Entrada: {ra['hora_entrada'].strftime('%d/%m/%Y %H:%M')}")

    st.markdown("### 🚘 Mis vehículos")
    vhs = obtener_vehiculos_de_usuario(user['id'])

    if not vhs:
        st.info("Aún no tienes vehículos registrados.")
    else:
        for v in vhs:
            with st.container(border=True):
                i = "🚗" if v['tipo'] == 'Auto' else "🏍️"
                st.markdown(f"**{i} {v['tipo']} — {v['placas']}**")
                st.caption(f"Marca: {v['marca'] or 'N/A'} | Modelo: {v['modelo'] or 'N/A'} | Color: {v['color'] or 'N/A'}")
                c1, c2 = st.columns(2)
                with c1:
                    if st.button("🎫 QR", key=f"qr_{v['id']}", use_container_width=True):
                        with st.spinner("🎫 Generando QR..."):
                            qd = {
                                "id_usuario": user['id'], "usuario": user['usuario'],
                                "nombre": user['nombre_completo'], "id_estudiante": user['id_estudiante'],
                                "matricula": user['matricula'], "carrera": user['carrera'],
                                "grupo": user['grupo'], "id_vehiculo": v['id'],
                                "tipo": v['tipo'], "placas": v['placas']
                            }
                            st.session_state.qr_generado = {"imagen": generar_qr_imagen(qd),
                                                            "vehiculo": v, "datos": qd}
                        st.rerun()
                with c2:
                    st.button("🗑️ Borrar", key=f"del_{v['id']}", use_container_width=True,
                              on_click=cb_mostrar_confirm, args=(f"confirmar_elim_veh_{v['id']}",))
                if st.session_state.get(f"confirmar_elim_veh_{v['id']}", False):
                    with st.container(border=True):
                        st.error(f"🚨 ¿Eliminar **{v['tipo']} {v['placas']}**?")
                        cs, cn = st.columns(2)
                        with cs:
                            st.button("✅ Sí", key=f"si_del_veh_{v['id']}", type="primary",
                                      use_container_width=True,
                                      on_click=cb_eliminar_vehiculo,
                                      args=(v['id'], user['id'], v['tipo'], v['placas']))
                        with cn:
                            st.button("❌ No", key=f"no_del_veh_{v['id']}", use_container_width=True,
                                      on_click=cb_ocultar_confirm, args=(f"confirmar_elim_veh_{v['id']}",))

    with st.expander("📜 Mi historial de visitas"):
        tv, vm = contar_visitas_usuario(user['id'])
        hist = obtener_historial_usuario(user['id'], limite=100)
        comp = [h for h in hist if h['hora_salida']]
        ps = "N/A"
        if comp:
            ts = [(h['hora_salida'] - h['hora_entrada']).total_seconds() / 60 for h in comp]
            p = sum(ts) / len(ts)
            ps = f"{p:.0f} min" if p < 60 else f"{p/60:.1f} hrs"
        c1, c2, c3 = st.columns(3)
        with c1: st.metric("🚗 Visitas", tv)
        with c2: st.metric("📅 Mes", vm)
        with c3: st.metric("⏱️ Promedio", ps)
        st.markdown("---")
        filtro = st.radio("Filtrar:", ["Todas", "Dentro", "Fuera"], horizontal=True,
                          label_visibility="collapsed", key="filtro_historial")
        if filtro == "Dentro":
            hf = [h for h in hist if h['estado'] == 'DENTRO']
        elif filtro == "Fuera":
            hf = [h for h in hist if h['estado'] == 'FUERA']
        else:
            hf = hist
        if not hf:
            st.info("No hay visitas.")
        else:
            kl = "hist_lazy_count"
            if kl not in st.session_state: st.session_state[kl] = 10
            vis = hf[:st.session_state[kl]]
            st.caption(f"Mostrando **{len(vis)}** de **{len(hf)}**")
            for h in vis:
                fe = h['hora_entrada'].strftime("%d/%m/%Y %H:%M") if h['hora_entrada'] else "N/A"
                fs = h['hora_salida'].strftime("%d/%m/%Y %H:%M") if h['hora_salida'] else None
                i = "🚗" if h['tipo'] == 'Auto' else "🏍️"
                ec = "#00ff88" if h['estado'] == 'DENTRO' else "#A89968"
                et = "🟢 DENTRO" if h['estado'] == 'DENTRO' else "🔴 COMPLETADA"
                if fs:
                    du = h['hora_salida'] - h['hora_entrada']
                    m = int(du.total_seconds() / 60)
                    dt = f"{m} min" if m < 60 else f"{m // 60}h {m % 60}min"
                else:
                    d = datetime.now() - h['hora_entrada']
                    m = int(d.total_seconds() / 60)
                    dt = f"{m} min (en curso)" if m < 60 else f"{m // 60}h {m % 60}min (en curso)"
                st.markdown(f"""
                    <div class="hist-item" style="border-left-color: {ec};">
                        <div style="display:flex; justify-content:space-between; flex-wrap:wrap;">
                            <span style="color:{ec}; font-weight:700; font-size:.85rem;">{et}</span>
                            <span class="hist-duracion">⏱️ {dt}</span>
                        </div>
                        <div style="color:#fff; margin-top:4px; font-size:.9rem;">{i} <b>{h['placas']}</b></div>
                        <div style="color:#aaa; font-size:.8rem;">⬇️ {fe}</div>
                        {f'<div style="color:#aaa; font-size:.8rem;">⬆️ {fs}</div>' if fs else ''}
                    </div>
                """, unsafe_allow_html=True)
            if st.session_state[kl] < len(hf):
                if st.button("⬇️ Ver más", use_container_width=True, key="hist_ver_mas"):
                    st.session_state[kl] += 10
                    st.rerun()

    with st.expander("➕ Registrar nuevo vehículo"):
        st.markdown("**Tipo de vehículo**")
        tipo = st.selectbox("Tipo", ["Auto", "Moto"], key="vh_tipo", label_visibility="collapsed")
        placas = st.text_input("Placas", key="vh_placas", placeholder="Ej. ABC-1234").upper().strip()
        pk = mostrar_validacion(placas, validar_placas, obligatorio=True)
        marca = st.text_input("Marca (opcional)", key="vh_marca")
        modelo = st.text_input("Modelo (opcional)", key="vh_modelo")
        color = st.text_input("Color (opcional)", key="vh_color")
        pu = True
        if placas and pk:
            plt = placas.replace("-", "").replace(" ", "").upper()
            if placas_existen(plt):
                st.markdown('<div class="val-error">❌ Ya existe un vehículo con esas placas</div>', unsafe_allow_html=True)
                pu = False
        tok = pk and pu
        if st.button("Registrar vehículo", use_container_width=True, type="primary",
                     disabled=not tok, key="vh_btn"):
            pl = placas.replace("-", "").replace(" ", "").upper()
            try:
                with st.spinner("💾 Registrando..."):
                    crear_vehiculo(user['id'], tipo, pl,
                                   marca.strip() if marca else None,
                                   modelo.strip() if modelo else None,
                                   color.strip() if color else None)
                    registrar_log(user['id'], "CREAR_VEHICULO",
                                  f"Vehículo {tipo} {pl} registrado", "Super_Vehiculos")
                limpiar_campos(['vh_placas', 'vh_marca', 'vh_modelo', 'vh_color'])
                set_flash("success", f"✅ Vehículo {pl} registrado.")
                st.rerun()
            except Exception as e:
                st.error(f"Error: {e}")

    if st.session_state.qr_generado:
        qi = st.session_state.qr_generado
        st.markdown("---")
        st.markdown("### 🎫 Tu código QR")
        st.info("Presenta este código en la caseta al entrar y salir.")
        _, cq, _ = st.columns([1, 2, 1])
        with cq:
            st.image(qi['imagen'], caption=f"QR — {qi['vehiculo']['tipo']} {qi['vehiculo']['placas']}")
        try:
            with st.spinner("📄 Generando PDF..."):
                pdf_bytes = generar_pdf_qr(user, qi['vehiculo'], qi['imagen'])
            st.download_button("📄 Descargar PDF", data=pdf_bytes,
                file_name=f"CUYPARK_{qi['vehiculo']['placas']}_{user['matricula'] or 'alumno'}.pdf",
                mime="application/pdf", use_container_width=True, type="primary")
        except Exception as e:
            st.error(f"Error al generar PDF: {e}")
        if st.button("❌ Cerrar QR", use_container_width=True):
            st.session_state.qr_generado = None
            st.rerun()


# ============================================================
# PANEL TRABAJADOR
# ============================================================
def panel_trabajador():
    user = st.session_state.usuario
    st.markdown(f'<div class="panel-header">👷 Caseta — {user["nombre_completo"]}</div>', unsafe_allow_html=True)
    mostrar_flash()
    mostrar_mi_cuenta(user)
    _contadores_caseta()
    st.caption("🟢 Contadores cada 15s")
    st.markdown("---")

    seccion = st.radio("Sección:", ["📷 Escanear QR", "📋 Vehículos Dentro"],
                       horizontal=True, label_visibility="collapsed")

    if seccion == "📷 Escanear QR":
        _seccion_escaner_qr(user)
    else:
        st.markdown("### 🚘 Vehículos dentro")
        ar = st.toggle("🔄 Auto-actualizar cada 10 segundos", value=False,
                       key="auto_refresh_caseta")
        sec = any(k.startswith("salida_rapida_") and v for k, v in st.session_state.items())
        if ar and not sec:
            st.caption("🟢 Actualizando cada 10 segundos")
            _render_dentro_auto(user)
        elif ar and sec:
            st.caption("⏸️ Auto-actualización pausada")
            _render_dentro_manual(user)
        else:
            if st.button("🔄 Refrescar ahora", use_container_width=True):
                st.rerun()
            _render_dentro_manual(user)


# ============================================================
# PANEL ADMIN
# ============================================================
def panel_admin():
    user = st.session_state.usuario
    st.markdown(f'<div class="panel-header">👑 Admin — {user["nombre_completo"]}</div>', unsafe_allow_html=True)
    mostrar_flash()
    _verificar_alertas_programadas(user)

    seccion = st.radio(
        "Sección:",
        ["📊 Dashboard", "👥 Usuarios", "📋 Registros", "📈 Métricas",
         "🔍 Auditoría", "💬 Mensajes", "🚗 Vehículos", "🔧 Mi Cuenta"],
        horizontal=True, label_visibility="collapsed"
    )

    if seccion == "📊 Dashboard":
        _dashboard_datos_vivo()

    elif seccion == "👥 Usuarios":
        st.markdown("### 👥 Gestión de Usuarios")
        sub = st.radio("Acción:", ["🎓 Crear Alumno", "👷 Crear Trabajador", "👑 Crear Admin", "📋 Ver Todos"],
                       horizontal=True, label_visibility="collapsed")

        if sub == "🎓 Crear Alumno":
            ids = generar_siguiente_id("ALU")
            st.info(f"💡 ID sugerido: **{ids}**")
            u = st.text_input("Usuario", key="ca_u", placeholder="mín. 3 caracteres, sin espacios")
            u_ok = mostrar_validacion(u, validar_usuario)
            p = st.text_input("Contraseña", type="password", key="ca_p")
            p_ok = mostrar_validacion(p, validar_password)
            nombre = st.text_input("Nombre completo", key="ca_n")
            n_ok = mostrar_validacion(nombre, validar_nombre)
            id_est = st.text_input("ID Estudiante", value=ids, key="ca_id")
            i_ok = mostrar_validacion(id_est, validar_id_estudiante)
            ca, cb = st.columns(2)
            with ca:
                mat = st.text_input("Matrícula (opcional)", key="ca_mat")
                m_ok = mostrar_validacion(mat, validar_matricula, obligatorio=False)
                car = st.text_input("Carrera (opcional)", key="ca_car")
                c_ok = mostrar_validacion(car, validar_carrera, obligatorio=False)
            with cb:
                gru = st.text_input("Grupo (opcional)", key="ca_gru")
                g_ok = mostrar_validacion(gru, validar_grupo, obligatorio=False)
                tel = st.text_input("Teléfono (opcional)", key="ca_tel")
                t_ok = mostrar_validacion(tel, validar_telefono, obligatorio=False)
            uu = True
            if u and u_ok and usuario_existe(u.lower().strip()):
                st.markdown('<div class="val-error">❌ Ese usuario ya existe</div>', unsafe_allow_html=True)
                uu = False
            iu = True
            if id_est and i_ok and id_estudiante_existe(id_est.strip()):
                st.markdown('<div class="val-error">❌ Ese ID ya está en uso</div>', unsafe_allow_html=True)
                iu = False
            tok = u_ok and p_ok and n_ok and i_ok and m_ok and c_ok and g_ok and t_ok and uu and iu
            if st.button("✅ Crear Alumno", use_container_width=True, type="primary",
                         disabled=not tok, key="ca_btn"):
                try:
                    with st.spinner("💾 Creando..."):
                        crear_usuario(usuario=u.lower().strip(), password=p, rol='alumno',
                            tipo_usuario='alumno', nombre_completo=nombre.strip(),
                            matricula=mat.strip() if mat else None,
                            carrera=car.strip() if car else None,
                            grupo=gru.strip() if gru else None,
                            telefono=tel.strip() if tel else None,
                            id_estudiante=id_est.strip() if id_est else None)
                        registrar_log(user['id'], "CREAR_ALUMNO",
                                      f"Alumno @{u.lower().strip()} ({nombre}) creado", "Super_Usuarios")
                    limpiar_campos(['ca_u', 'ca_p', 'ca_n', 'ca_id', 'ca_mat', 'ca_car', 'ca_gru', 'ca_tel'])
                    set_flash("success", f"✅ Alumno **{nombre}** creado.")
                    st.rerun()
                except Exception as e:
                    st.error(f"Error: {e}")

        elif sub == "👷 Crear Trabajador":
            u = st.text_input("Usuario", key="ct_u")
            u_ok = mostrar_validacion(u, validar_usuario)
            p = st.text_input("Contraseña", type="password", key="ct_p")
            p_ok = mostrar_validacion(p, validar_password)
            nombre = st.text_input("Nombre completo", key="ct_n")
            n_ok = mostrar_validacion(nombre, validar_nombre)
            tel = st.text_input("Teléfono (opcional)", key="ct_tel")
            t_ok = mostrar_validacion(tel, validar_telefono, obligatorio=False)
            uu = True
            if u and u_ok and usuario_existe(u.lower().strip()):
                st.markdown('<div class="val-error">❌ Ese usuario ya existe</div>', unsafe_allow_html=True)
                uu = False
            tok = u_ok and p_ok and n_ok and t_ok and uu
            if st.button("✅ Crear Trabajador", use_container_width=True, type="primary",
                         disabled=not tok, key="ct_btn"):
                try:
                    with st.spinner("💾 Creando..."):
                        crear_usuario(usuario=u.lower().strip(), password=p, rol='trabajador',
                            tipo_usuario='administrativo', nombre_completo=nombre.strip(),
                            telefono=tel.strip() if tel else None)
                        registrar_log(user['id'], "CREAR_TRABAJADOR",
                                      f"Trabajador @{u.lower().strip()} ({nombre}) creado", "Super_Usuarios")
                    limpiar_campos(['ct_u', 'ct_p', 'ct_n', 'ct_tel'])
                    set_flash("success", f"✅ Trabajador **{nombre}** creado.")
                    st.rerun()
                except Exception as e:
                    st.error(f"Error: {e}")

        elif sub == "👑 Crear Admin":
            st.warning("⚠️ Los admins tienen acceso total.")
            u = st.text_input("Usuario", key="cA_u")
            u_ok = mostrar_validacion(u, validar_usuario)
            p = st.text_input("Contraseña", type="password", key="cA_p")
            p_ok = mostrar_validacion(p, validar_password)
            nombre = st.text_input("Nombre completo", key="cA_n")
            n_ok = mostrar_validacion(nombre, validar_nombre)
            tel = st.text_input("Teléfono (opcional)", key="cA_tel")
            t_ok = mostrar_validacion(tel, validar_telefono, obligatorio=False)
            uu = True
            if u and u_ok and usuario_existe(u.lower().strip()):
                st.markdown('<div class="val-error">❌ Ese usuario ya existe</div>', unsafe_allow_html=True)
                uu = False
            tok = u_ok and p_ok and n_ok and t_ok and uu
            if st.button("✅ Crear Administrador", use_container_width=True, type="primary",
                         disabled=not tok, key="cA_btn"):
                try:
                    with st.spinner("💾 Creando..."):
                        crear_usuario(usuario=u.lower().strip(), password=p, rol='admin',
                            tipo_usuario='administrativo', nombre_completo=nombre.strip(),
                            telefono=tel.strip() if tel else None)
                        registrar_log(user['id'], "CREAR_ADMIN",
                                      f"Administrador @{u.lower().strip()} ({nombre}) creado", "Super_Usuarios")
                    limpiar_campos(['cA_u', 'cA_p', 'cA_n', 'cA_tel'])
                    set_flash("success", f"✅ Administrador **{nombre}** creado.")
                    st.rerun()
                except Exception as e:
                    st.error(f"Error: {e}")

        elif sub == "📋 Ver Todos":
            cf1, cf2 = st.columns([2, 1])
            with cf1:
                filtro = st.selectbox("Filtrar por rol", ["Todos", "alumno", "trabajador", "admin"])
            with cf2:
                sa = st.checkbox("Solo activos", value=False)
            with st.spinner("📋 Cargando..."):
                us = obtener_todos_usuarios(None if filtro == "Todos" else filtro, solo_activos=sa)
            if not us:
                st.info("No hay usuarios.")
            else:
                ac = sum(1 for u in us if u['activo'])
                ic = len(us) - ac
                st.write(f"**Total: {len(us)}** ({ac} activos, {ic} inactivos)")

                with st.expander("📊 Comparativa activos vs inactivos"):
                    datos = contar_usuarios_por_estado()
                    dfc = pd.DataFrame([
                        {"Rol": r.capitalize(), "Activos": d["activos"], "Inactivos": d["inactivos"]}
                        for r, d in datos.items()
                    ])
                    st.dataframe(dfc, use_container_width=True, hide_index=True)
                    fc = go.Figure(data=[
                        go.Bar(name='Activos', x=dfc['Rol'], y=dfc['Activos'],
                               marker_color='#00ff88', text=dfc['Activos'], textposition='outside'),
                        go.Bar(name='Inactivos', x=dfc['Rol'], y=dfc['Inactivos'],
                               marker_color='#D7192D', text=dfc['Inactivos'], textposition='outside'),
                    ])
                    fc.update_layout(
                        barmode='group', paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
                        font=dict(color=COLOR_CREMA, size=11),
                        xaxis=dict(gridcolor="rgba(201,169,97,0.1)"),
                        yaxis=dict(gridcolor="rgba(201,169,97,0.1)"),
                        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
                        margin=dict(l=10, r=10, t=40, b=10), height=300,
                    )
                    st.plotly_chart(fc, use_container_width=True, config={"displayModeBar": False})
                    if st.button("⚡ Crear índices recomendados", use_container_width=True, key="btn_indices"):
                        with st.spinner("Creando..."):
                            cr, er = crear_indices()
                        st.success(f"✅ {len(cr)} índice(s) creado(s).")
                        if er: st.caption(f"ℹ️ {len(er)} ya existían.")

                st.markdown("---")
                for u in us:
                    with st.container(border=True):
                        icono = {'alumno': '🎓', 'trabajador': '👷', 'admin': '👑'}.get(u['rol'], '👤')
                        eb = "🟢" if u['activo'] else "🔒"
                        c1, c2, c3, c4, c5 = st.columns([4, 1, 1, 1, 1])
                        with c1:
                            ba = False; mb = 0
                            if u.get('bloqueado_hasta') and isinstance(u['bloqueado_hasta'], datetime):
                                if u['bloqueado_hasta'] > datetime.now():
                                    ba = True
                                    mb = int((u['bloqueado_hasta'] - datetime.now()).total_seconds() / 60) + 1
                            if ba: eb = "🔐"
                            st.markdown(f"**{eb} {icono} {u['nombre_completo']}** — @{u['usuario']}")
                            if not u['activo']: st.caption("🔒 Desactivada")
                            if ba: st.caption(f"🔐 Bloqueada — {mb} min")
                            i_f = u.get('intentos_fallidos') or 0
                            if i_f > 0 and not ba: st.caption(f"⚠️ Intentos fallidos: {i_f}/5")
                            if u['id_estudiante']: st.caption(f"ID: {u['id_estudiante']}")
                            st.caption(f"Rol: {u['rol']} | Tel: {u['telefono'] or 'N/A'}")
                            if u['rol'] == 'alumno':
                                st.caption(f"Matrícula: {u['matricula'] or 'N/A'}")
                        with c2:
                            if st.button("✏️", key=f"edit_{u['id']}", use_container_width=True):
                                for o in us:
                                    if o['id'] != u['id']:
                                        st.session_state[f"editando_{o['id']}"] = False
                                st.session_state[f"editando_{u['id']}"] = not st.session_state.get(f"editando_{u['id']}", False)
                                st.rerun()
                        with c3:
                            if u['id'] == user['id']:
                                st.caption("(Tú)")
                            else:
                                if u['activo']:
                                    st.button("🔒", key=f"lock_{u['id']}", use_container_width=True,
                                              on_click=cb_mostrar_confirm, args=(f"confirmar_desactivar_{u['id']}",))
                                else:
                                    st.button("🔓", key=f"unlock_{u['id']}", use_container_width=True,
                                              on_click=cb_activar_usuario,
                                              args=(u['id'], user['id'], user['usuario'], u['usuario'], u['nombre_completo']))
                        with c4:
                            if u['id'] != user['id']:
                                st.button("🗑️", key=f"del_user_{u['id']}", use_container_width=True,
                                          on_click=cb_mostrar_confirm, args=(f"confirmar_eliminar_{u['id']}",))
                        with c5:
                            bb = False
                            if u.get('bloqueado_hasta') and isinstance(u['bloqueado_hasta'], datetime):
                                if u['bloqueado_hasta'] > datetime.now():
                                    bb = True
                            if u['id'] == user['id']:
                                st.caption("")
                            elif bb:
                                if st.button("🔐", key=f"unlock_bloq_{u['id']}", use_container_width=True):
                                    ok, msg = desbloquear_usuario(u['id'])
                                    if ok:
                                        registrar_log(user['id'], "DESBLOQUEAR_USUARIO",
                                                      f"Cuenta @{u['usuario']} desbloqueada", "Super_Usuarios", u['id'])
                                        set_flash("success", msg)
                                    else:
                                        set_flash("error", msg)
                                    st.rerun()

                        if st.session_state.get(f"confirmar_eliminar_{u['id']}", False):
                            with st.container(border=True):
                                st.error(f"🚨 ¿Eliminar **{u['nombre_completo']}**?")
                                if u['rol'] == 'admin':
                                    st.warning("⚠️ Es ADMINISTRADOR.")
                                    tx = st.text_input("Escribe ELIMINAR:", key=f"confirma_texto_{u['id']}")
                                    cok = (tx.strip() == "ELIMINAR")
                                else:
                                    cok = True
                                cs, cn = st.columns(2)
                                with cs:
                                    st.button("✅ Sí", key=f"si_del_{u['id']}", type="primary",
                                              use_container_width=True, disabled=not cok,
                                              on_click=cb_eliminar_usuario,
                                              args=(u['id'], user['id'], user['usuario'], u['usuario'], u['nombre_completo']))
                                with cn:
                                    st.button("❌ No", key=f"no_del_{u['id']}", use_container_width=True,
                                              on_click=cb_ocultar_confirm, args=(f"confirmar_eliminar_{u['id']}",))

                        if st.session_state.get(f"confirmar_desactivar_{u['id']}", False):
                            with st.container(border=True):
                                st.warning(f"🔒 ¿Desactivar **{u['nombre_completo']}**?")
                                cs, cn = st.columns(2)
                                with cs:
                                    st.button("✅ Sí", key=f"si_desc_{u['id']}", type="primary",
                                              use_container_width=True,
                                              on_click=cb_desactivar_usuario,
                                              args=(u['id'], user['id'], user['usuario'], u['usuario'], u['nombre_completo']))
                                with cn:
                                    st.button("❌ No", key=f"no_desc_{u['id']}", use_container_width=True,
                                              on_click=cb_ocultar_confirm, args=(f"confirmar_desactivar_{u['id']}",))

                    if st.session_state.get(f"editando_{u['id']}", False):
                        st.markdown('<div class="edit-form">', unsafe_allow_html=True)
                        st.markdown(f"#### ✏️ Editando: @{u['usuario']}")
                        en = st.text_input("Nombre completo", value=u['nombre_completo'], key=f"ed_n_{u['id']}")
                        en_ok = mostrar_validacion(en, validar_nombre)
                        et = st.text_input("Teléfono", value=u['telefono'] or "", key=f"ed_tel_{u['id']}")
                        et_ok = mostrar_validacion(et, validar_telefono, obligatorio=False)
                        if u['rol'] == 'alumno':
                            ea, eb = st.columns(2)
                            with ea:
                                eid = st.text_input("ID", value=u['id_estudiante'] or "", key=f"ed_id_{u['id']}")
                                eid_ok = mostrar_validacion(eid, validar_id_estudiante, obligatorio=False)
                                em = st.text_input("Matrícula", value=u['matricula'] or "", key=f"ed_mat_{u['id']}")
                                em_ok = mostrar_validacion(em, validar_matricula, obligatorio=False)
                            with eb:
                                ec = st.text_input("Carrera", value=u['carrera'] or "", key=f"ed_car_{u['id']}")
                                ec_ok = mostrar_validacion(ec, validar_carrera, obligatorio=False)
                                eg = st.text_input("Grupo", value=u['grupo'] or "", key=f"ed_gru_{u['id']}")
                                eg_ok = mostrar_validacion(eg, validar_grupo, obligatorio=False)
                            eiu = True
                            if eid and eid_ok and id_estudiante_existe_otro(eid.strip(), u['id']):
                                st.markdown('<div class="val-error">❌ ID ya lo usa otro</div>', unsafe_allow_html=True)
                                eiu = False
                        else:
                            eid = u['id_estudiante']; eid_ok = True; eiu = True
                            em = u['matricula']; em_ok = True
                            ec = u['carrera']; ec_ok = True
                            eg = u['grupo']; eg_ok = True
                        st.markdown("##### 🔐 Cambiar contraseña (opcional)")
                        ep = st.text_input("Nueva contraseña", type="password", key=f"ed_p_{u['id']}")
                        ep2 = st.text_input("Confirmar", type="password", key=f"ed_p2_{u['id']}")
                        pok = True; pm = ""
                        if ep:
                            if len(ep) < 3: pok = False; pm = "Mínimo 3"
                            elif ep != ep2: pok = False; pm = "No coinciden"
                        if ep:
                            if pok: st.markdown('<div class="val-ok">✅ Válida</div>', unsafe_allow_html=True)
                            else: st.markdown(f'<div class="val-error">❌ {pm}</div>', unsafe_allow_html=True)
                        tok = en_ok and et_ok and eid_ok and em_ok and ec_ok and eg_ok and eiu and pok
                        cx, cy = st.columns(2)
                        with cx:
                            if st.button("💾 Guardar", use_container_width=True, type="primary",
                                         disabled=not tok, key=f"ed_save_{u['id']}"):
                                try:
                                    with st.spinner("💾 Guardando..."):
                                        actualizar_usuario(
                                            id_usuario=u['id'], nombre_completo=en.strip(),
                                            telefono=et.strip() or None,
                                            matricula=em.strip() or None,
                                            carrera=ec.strip() or None,
                                            grupo=eg.strip() or None,
                                            id_estudiante=eid.strip() or None,
                                            tipo_usuario=u['tipo_usuario'],
                                            nueva_password=ep if ep else None)
                                        registrar_log(user['id'], "EDITAR_USUARIO",
                                                      f"Usuario @{u['usuario']} editado", "Super_Usuarios", u['id'])
                                    st.session_state[f"editando_{u['id']}"] = False
                                    limpiar_campos([f"ed_n_{u['id']}", f"ed_tel_{u['id']}", f"ed_id_{u['id']}",
                                                    f"ed_mat_{u['id']}", f"ed_car_{u['id']}", f"ed_gru_{u['id']}",
                                                    f"ed_p_{u['id']}", f"ed_p2_{u['id']}"])
                                    set_flash("success", "✅ Usuario actualizado.")
                                    st.rerun()
                                except Exception as e:
                                    st.error(f"Error: {e}")
                        with cy:
                            if st.button("❌ Cancelar", use_container_width=True, key=f"ed_cancel_{u['id']}"):
                                st.session_state[f"editando_{u['id']}"] = False
                                limpiar_campos([f"ed_n_{u['id']}", f"ed_tel_{u['id']}", f"ed_id_{u['id']}",
                                                f"ed_mat_{u['id']}", f"ed_car_{u['id']}", f"ed_gru_{u['id']}",
                                                f"ed_p_{u['id']}", f"ed_p2_{u['id']}"])
                                st.rerun()
                        st.markdown('</div>', unsafe_allow_html=True)

    elif seccion == "📋 Registros":
        st.markdown("### 📋 Registros de entradas/salidas")
        fd, fh = selector_rango_fechas("registros")

        c1, c2 = st.columns(2)
        with c1: fe = st.selectbox("Estado", ["Todos", "DENTRO", "FUERA"], key="reg_estado")
        with c2: ft = st.selectbox("Tipo", ["Todos", "Auto", "Moto"], key="reg_tipo")

        buscar = st.text_input("🔍 Buscar (nombre, placas, matrícula)", key="reg_buscar").lower().strip()
        pp = mostrar_paginacion_superior("reg")

        fa = f"{fd}|{fh}|{fe}|{ft}|{buscar}|{pp}"
        if st.session_state.get("reg_filtros_prev") != fa:
            st.session_state["reg_pagina"] = 1
            st.session_state["reg_filtros_prev"] = fa

        pag = st.session_state.get("reg_pagina", 1)
        off = (pag - 1) * pp

        total = contar_registros_filtrados(fd, fh, fe, ft, buscar or None)
        regs = obtener_registros_paginado(fd, fh, fe, ft, buscar or None, off, pp)

        if not regs:
            st.info("No hay registros en el rango seleccionado.")
        else:
            st.write(f"**Total de registros: {total}**")

            df_exp = pd.DataFrame([{
                'ID': r['id'],
                'Estado': r['estado'],
                'Alumno': r['nombre_completo'],
                'Matrícula': r['matricula'] or '',
                'Carrera': r['carrera'] or '',
                'Tipo': r['tipo'],
                'Placas': r['placas'],
                'Fecha Entrada': r['hora_entrada'].strftime("%d/%m/%Y") if r['hora_entrada'] else '',
                'Hora Entrada': r['hora_entrada'].strftime("%H:%M:%S") if r['hora_entrada'] else '',
                'Fecha Salida': r['hora_salida'].strftime("%d/%m/%Y") if r['hora_salida'] else 'En curso',
                'Hora Salida': r['hora_salida'].strftime("%H:%M:%S") if r['hora_salida'] else ''
            } for r in regs])

            rango_txt = (f"Período: {fd.strftime('%d/%m/%Y')} - {fh.strftime('%d/%m/%Y')}"
                         if fd and fh else "Período: Todo el historial")

            pdf_bytes = None
            try:
                pdf_bytes = generar_pdf_reporte(
                    df_exp, "Reporte de Registros de Estacionamiento",
                    f"{rango_txt} | Total: {total} registro(s)"
                )
            except Exception as e:
                st.error(f"Error generando PDF: {e}")

            if pdf_bytes:
                st.download_button(
                    "📄 Descargar Reporte PDF",
                    data=pdf_bytes,
                    file_name=f"registros_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf",
                    mime="application/pdf",
                    use_container_width=True,
                    type="primary",
                    key="dl_reg_pdf"
                )

            with st.expander("📸 Ver evidencias fotográficas"):
                ids_con_ev = [r['id'] for r in regs if r['evidencia_entrada'] or r['evidencia_salida']]
                if not ids_con_ev:
                    st.info("Ninguno de los registros mostrados tiene evidencias.")
                else:
                    sel_id = st.selectbox("Selecciona un ID de registro", ids_con_ev, key="ev_sel_reg")
                    reg = next((x for x in regs if x['id'] == sel_id), None)
                    if reg:
                        st.markdown(f"**Registro #{reg['id']} — {reg['placas']} ({reg['nombre_completo']})**")
                        if reg['evidencia_entrada']:
                            st.markdown("**📷 Entrada:**")
                            img = base64_a_bytes(reg['evidencia_entrada'])
                            if img:
                                st.image(img, use_container_width=True)
                        if reg['evidencia_salida']:
                            st.markdown("**📷 Salida:**")
                            img = base64_a_bytes(reg['evidencia_salida'])
                            if img:
                                st.image(img, use_container_width=True)

            st.markdown("---")
            render_paginacion_inferior("reg", pag, total, pp, "reg_pagina")

    elif seccion == "📈 Métricas":
        st.markdown("### 📈 Métricas y patrones")
        with st.spinner("📈 Calculando..."):
            regs = obtener_todos_los_registros()
        if not regs:
            st.info("Sin datos suficientes.")
        else:
            df = pd.DataFrame(regs)
            df['hora_entrada'] = pd.to_datetime(df['hora_entrada'], errors='coerce')
            df['hora_salida'] = pd.to_datetime(df['hora_salida'], errors='coerce')
            st.markdown("#### ⏰ Horas de mayor demanda")
            df['hora_del_dia'] = df['hora_entrada'].dt.hour
            hp = df.groupby('hora_del_dia').size().reset_index(name='entradas').sort_values('hora_del_dia')
            st.bar_chart(hp.set_index('hora_del_dia')['entradas'])
            st.markdown("#### ⏱️ Permanencia promedio")
            comp = df.dropna(subset=['hora_salida']).copy()
            if not comp.empty:
                comp['dur'] = (comp['hora_salida'] - comp['hora_entrada']).dt.total_seconds() / 60
                st.metric("Promedio general", f"{comp['dur'].mean():.1f} min")
                c1, c2 = st.columns(2)
                with c1:
                    ap = comp[comp['tipo'] == 'Auto']['dur'].mean()
                    if pd.notna(ap): st.metric("🚗 Autos", f"{ap:.1f} min")
                with c2:
                    mp = comp[comp['tipo'] == 'Moto']['dur'].mean()
                    if pd.notna(mp): st.metric("🏍️ Motos", f"{mp:.1f} min")
            st.markdown("#### 🎓 Uso por carrera")
            pc = df[df['carrera'].notna()].groupby('carrera').size().reset_index(name='usos')
            if not pc.empty: st.bar_chart(pc.set_index('carrera')['usos'])
            st.markdown("#### 🚗 Uso por tipo")
            pt = df.groupby('tipo').size().reset_index(name='cantidad')
            st.dataframe(pt, use_container_width=True, hide_index=True)

    elif seccion == "🔍 Auditoría":
        st.markdown("### 🔍 Registro de Auditoría")
        st.caption("Historial de acciones importantes.")
        st.metric("Total de registros", contar_logs())

        fd, fh = selector_rango_fechas("auditoria")

        c1, c2 = st.columns(2)
        with c1:
            acciones = ["Todas"] + obtener_acciones_unicas()
            fac = st.selectbox("Filtrar por acción", acciones, key="aud_accion")
        with c2:
            buscar = st.text_input("🔍 Buscar (usuario o detalle)", key="aud_buscar").lower().strip()

        pp = mostrar_paginacion_superior("aud")

        fa = f"{fd}|{fh}|{fac}|{buscar}|{pp}"
        if st.session_state.get("aud_filtros_prev") != fa:
            st.session_state["aud_pagina"] = 1
            st.session_state["aud_filtros_prev"] = fa

        pag = st.session_state.get("aud_pagina", 1)
        off = (pag - 1) * pp

        total = contar_logs_filtrados(fac, buscar or None, fd, fh)
        logs = obtener_logs(limite=pp, filtro_accion=fac,
                            buscar=buscar if buscar else None,
                            fecha_desde=fd, fecha_hasta=fh, offset=off)

        if not logs:
            st.info("No hay logs que coincidan con los filtros.")
        else:
            st.write(f"**Total de eventos: {total}**")

            df_logs = pd.DataFrame([{
                'ID': l['id'],
                'Fecha': l['fecha'].strftime("%d/%m/%Y %H:%M:%S") if l['fecha'] else '',
                'Acción': l['accion'],
                'Usuario': f"@{l['usuario_accion']}" if l['usuario_accion'] else "Sistema",
                'Nombre': l['nombre_accion'] or '',
                'Rol': l['rol_accion'] or '',
                'Detalles': l['detalles'] or '',
                'Tabla afectada': l['tabla_afectada'] or '',
                'ID afectado': l['id_afectado'] if l['id_afectado'] else ''
            } for l in logs])

            rango_txt = (f"Período: {fd.strftime('%d/%m/%Y')} - {fh.strftime('%d/%m/%Y')}"
                         if fd and fh else "Período: Todo el historial")

            pdf_bytes = None
            try:
                pdf_bytes = generar_pdf_reporte(
                    df_logs, "Reporte de Auditoría del Sistema",
                    f"{rango_txt} | Total: {total} evento(s)"
                )
            except Exception as e:
                st.error(f"Error generando PDF: {e}")

            if pdf_bytes:
                st.download_button(
                    "📄 Descargar Reporte PDF",
                    data=pdf_bytes,
                    file_name=f"auditoria_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf",
                    mime="application/pdf",
                    use_container_width=True,
                    type="primary",
                    key="dl_aud_pdf"
                )

            st.markdown("---")
            render_paginacion_inferior("aud", pag, total, pp, "aud_pagina")

    elif seccion == "💬 Mensajes":
        st.markdown("### 💬 Mensajes")
        t1, t2, t3 = st.tabs(["📥 Recibidos", "📤 Enviados", "✉️ Nuevo"])
        with t1:
            ms = obtener_mensajes_para_usuario(user['id'])
            if not ms:
                st.info("No tienes mensajes.")
            else:
                nl = sum(1 for m in ms if not m['leido'])
                st.caption(f"**{nl}** sin leer de **{len(ms)}**")
                for m in ms:
                    il = "📭" if m['leido'] else "📬"
                    pr = "🔵 " if not m['leido'] else ""
                    ti = {"mensaje": "💬", "alerta": "🚨", "aviso": "📢"}.get(m['tipo'], "💬")
                    with st.container(border=True):
                        c1, c2 = st.columns([5, 1])
                        with c1:
                            st.markdown(f"**{pr}{ti} {m['asunto']}**")
                            st.caption(f"De: @{m['remitente_usuario']} · {m['fecha'].strftime('%d/%m/%Y %H:%M')}")
                            if m.get('es_broadcast'): st.caption("📢 General")
                        with c2:
                            if not m['leido']:
                                if st.button("✅", key=f"leido_{m['id']}", use_container_width=True):
                                    marcar_mensaje_leido(m['id']); st.rerun()
                        with st.expander("Ver contenido"):
                            st.write(m['cuerpo'])
                            if st.button("🗑️ Eliminar", key=f"del_msg_{m['id']}", use_container_width=True):
                                eliminar_mensaje(m['id']); st.rerun()
        with t2:
            env = obtener_mensajes_enviados(user['id'])
            if not env:
                st.info("No has enviado mensajes.")
            else:
                for m in env:
                    with st.container(border=True):
                        st.markdown(f"**{m['asunto']}**")
                        d = "📢 Todos los alumnos" if m.get('es_broadcast') else f"@{m['destinatario_usuario']}"
                        st.caption(f"Para: {d} · {m['fecha'].strftime('%d/%m/%Y %H:%M')}")
                        with st.expander("Ver contenido"):
                            st.write(m['cuerpo'])
                            if st.button("🗑️ Eliminar", key=f"del_env_{m['id']}", use_container_width=True):
                                eliminar_mensaje(m['id']); st.rerun()
        with t3:
            st.markdown("#### ✉️ Enviar mensaje")
            dt = st.radio("Destinatario", ["📢 Todos los alumnos (broadcast)", "🎓 Un alumno",
                                            "👷 Un trabajador", "👑 Un admin"],
                          horizontal=False, key="msg_dest_tipo")
            did = None
            if "Todos" in dt:
                did = None
            else:
                if "alumno" in dt: ud = obtener_todos_usuarios("alumno", solo_activos=True)
                elif "trabajador" in dt: ud = obtener_todos_usuarios("trabajador", solo_activos=True)
                else: ud = obtener_todos_usuarios("admin", solo_activos=True)
                op = {f"{u['nombre_completo']} (@{u['usuario']})": u['id'] for u in ud}
                if op:
                    sel = st.selectbox("Selecciona", list(op.keys()), key="msg_dest_sel")
                    did = op[sel]
                else:
                    st.warning("No hay usuarios."); did = -1
            asunto = st.text_input("Asunto", key="msg_asunto", max_chars=200)
            tipo_msg = st.selectbox("Tipo", ["mensaje", "aviso", "alerta"], key="msg_tipo")
            cuerpo = st.text_area("Mensaje", key="msg_cuerpo", height=150)
            aok = bool(asunto and asunto.strip())
            cok = bool(cuerpo and cuerpo.strip())
            dok = did != -1
            if st.button("📤 Enviar", type="primary", use_container_width=True,
                         disabled=not (aok and cok and dok), key="msg_enviar"):
                enviar_mensaje(user['id'], asunto.strip(), cuerpo.strip(), did, tipo_msg)
                registrar_log(user['id'], "ENVIAR_MENSAJE",
                              f"@{user['usuario']} envió '{asunto[:60]}'", "Super_Mensajes")
                limpiar_campos(['msg_asunto', 'msg_cuerpo'])
                set_flash("success", "✅ Mensaje enviado.")
                st.rerun()

    elif seccion == "🚗 Vehículos":
        st.markdown("### 🚗 Vehículos registrados")
        st.caption("Busca un vehículo para ver su historial.")
        bv = st.text_input("🔍 Buscar por placas, alumno o matrícula", key="veh_buscar_admin").strip()
        with st.spinner("🔎 Buscando..."):
            vh = buscar_vehiculos_admin(bv if bv else None, limite=50)
        if not vh:
            st.info("No se encontraron vehículos.")
        else:
            st.caption(f"**{len(vh)}** vehículo(s)")
            for v in vh:
                i = "🚗" if v['tipo'] == 'Auto' else "🏍️"
                with st.container(border=True):
                    c1, c2 = st.columns([4, 1])
                    with c1:
                        st.markdown(f"**{i} {v['placas']}** — {v['tipo']}")
                        st.caption(f"Dueño: {v['nombre_completo']} (@{v['usuario']})")
                        st.caption(f"Matrícula: {v['matricula'] or 'N/A'} | ID: {v['id_estudiante'] or 'N/A'}")
                        st.caption(f"Marca: {v['marca'] or 'N/A'} | Modelo: {v['modelo'] or 'N/A'}")
                    with c2:
                        if st.button("📜 Historial", key=f"hist_veh_{v['id']}", use_container_width=True):
                            st.session_state[f"ver_hist_veh_{v['id']}"] = not st.session_state.get(f"ver_hist_veh_{v['id']}", False)
                            st.rerun()
                    if st.session_state.get(f"ver_hist_veh_{v['id']}", False):
                        st.markdown("---")
                        st.markdown(f"#### 📜 Historial de {v['placas']}")
                        hist = obtener_historial_vehiculo(v['id'])
                        if not hist:
                            st.info("Sin registros.")
                        else:
                            tv = len(hist)
                            comp = [h for h in hist if h['hora_salida']]
                            act = any(h['estado'] == 'DENTRO' for h in hist)
                            ca, cb, cc = st.columns(3)
                            with ca: st.metric("Visitas", tv)
                            with cb:
                                ps = "N/A"
                                if comp:
                                    ts = [(h['hora_salida'] - h['hora_entrada']).total_seconds() / 60 for h in comp]
                                    p = sum(ts) / len(ts)
                                    ps = f"{p:.0f} min" if p < 60 else f"{p/60:.1f} h"
                                st.metric("Permanencia prom.", ps)
                            with cc:
                                st.metric("Estado", "🟢 DENTRO" if act else "🔴 FUERA")
                            if st.button("❌ Cerrar historial", key=f"cerrar_hist_{v['id']}", use_container_width=True):
                                st.session_state[f"ver_hist_veh_{v['id']}"] = False
                                st.rerun()

    elif seccion == "🔧 Mi Cuenta":
        st.markdown("### 🔧 Mi Cuenta")
        with st.expander("👤 Ver mi información"):
            st.write(f"**Usuario:** {user['usuario']}")
            st.write(f"**Nombre:** {user['nombre_completo']}")
            st.write(f"**Rol:** {user['rol']}")
            st.write(f"**Tipo:** {user['tipo_usuario'] or 'N/A'}")
            st.write(f"**Teléfono:** {user['telefono'] or 'N/A'}")
            st.write(f"**Registrado:** {user['fecha_registro']}")
        mostrar_mi_cuenta(user)


# ============================================================
# MAIN
# ============================================================
if st.session_state.usuario is not None:
    rol = st.session_state.usuario['rol']
    iconos = {'alumno': '🎓', 'trabajador': '👷', 'admin': '👑'}
    c1, c2, c3 = st.columns([3, 2, 1])
    with c1:
        st.markdown(
            f"<div style='padding-top: 10px; color: #C9A961;'>"
            f"<b>{iconos.get(rol, '👤')} @{st.session_state.usuario['usuario']}</b>"
            f"</div>", unsafe_allow_html=True)
    with c2:
        st.markdown(mostrar_marca_cudy(), unsafe_allow_html=True)
    with c3:
        if st.button("🚪 Salir", use_container_width=True):
            cerrar_sesion()

if st.session_state.usuario is None:
    pantalla_login()
else:
    rol = st.session_state.usuario['rol']
    if rol == 'alumno': panel_alumno()
    elif rol == 'trabajador': panel_trabajador()
    elif rol == 'admin': panel_admin()
    else: st.error(f"Rol desconocido: {rol}")